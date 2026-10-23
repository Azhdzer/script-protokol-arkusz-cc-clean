# -*- coding: utf-8 -*-
"""
Przetworniki: tabela w panelu (PRZETWORNIKI), PZ i parametry kopii arkusza.

Zlecenie 230: dwa przetworniki wilgotnosci Galltec #FPC3/5 (4-20 mA), mierzone
multimetrem 1586A na kanalach 001 i 221. Kopia arkusza obliczeniowego miala
dane w E/F (tm/RHm) i parametry termohigrometru; poprawnie:
  * dane w G/H (Itm/IRHm), K10 'Przetwornik', L10 'Prąd' (jednostka 'mA'
    ze Strony 3), K14 '1586A', K16 '001' (kanal RH), M22/O22 = 4 ÷ 20,
  * zakres z opisu skalowania na Stronie 2, kanaly z Uwag — albo z tabeli
    'Przetworniki' w panelu (kanal t i kanal RH osobno).
PZ lezal w folderze PZ jako skrot .lnk, a blok 'UWAGA: 1) ÷ 6) ...' dawal
fikcyjne przyrzady.
"""

import io
import os
import sys
import unittest
from contextlib import redirect_stdout

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import generuj_arkusze as A
import generuj_obserwacje as G
import pz_dane

TABELA = [["263051", "-", "001", "1586A", "Prąd", "4", "20", "", "", "0", "100"],
          ["D0540007", "105", "115", "1586A-02", "Napięcie", "0", "5", "-40", "60", "0", "100"],
          ["", "", "", "", "", "", "", "", "", "", ""]]


class TestTabela(unittest.TestCase):

    def setUp(self):
        self.t = pz_dane.przetworniki_z_tabeli(TABELA)

    def test_puste_wiersze_pomijane(self):
        self.assertEqual(len(self.t), 2)

    def test_kanaly_t_i_rh_osobno(self):
        p = self.t[pz_dane.normalizuj_serial("D0540007")]
        self.assertEqual((p.kanal_t, p.kanal_rh), ("105", "115"))
        p = self.t[pz_dane.normalizuj_serial("263051")]
        self.assertEqual((p.kanal_t, p.kanal_rh), ("", "001"))

    def test_sygnal_i_jednostka(self):
        self.assertEqual(self.t["263051"].jednostka, "mA")
        self.assertEqual(self.t[pz_dane.normalizuj_serial("D0540007")].jednostka, "V")

    def test_opis_skalowania_jak_recznie(self):
        p = self.t["263051"]
        self.assertEqual(pz_dane.opis_skalowania(p, "rh"),
                         "(4 ÷ 20) mA\nodpowiada:\n(0 ÷ 100) %")
        self.assertIsNone(pz_dane.opis_skalowania(p, "t"))

    def test_sygnal_z_tekstu(self):
        for tekst, ocz in (("Prąd", "Prąd"), ("prad", "Prąd"), ("mA", "Prąd"), ("I", "Prąd"),
                           ("Napięcie", "Napięcie"), ("V", "Napięcie"), ("U", "Napięcie"),
                           ("", None)):
            with self.subTest(tekst=tekst):
                self.assertEqual(pz_dane.sygnal_z_tekstu(tekst), ocz)


class TestOpisyZProtokolu(unittest.TestCase):

    def test_skalowanie_4_20(self):
        self.assertEqual(pz_dane.zakresy_ze_skalowania("(4 ÷ 20) mA\nodpowiada:\n(0 ÷ 100) %"),
                         ((4.0, 20.0, "mA"), (0.0, 100.0)))

    def test_skalowanie_reczne_0_5V(self):
        self.assertEqual(pz_dane.zakresy_ze_skalowania("0-5V \n(-40 / +60)'C"),
                         ((0.0, 5.0, "V"), (-40.0, 60.0)))

    def test_skalowanie_brak(self):
        self.assertIsNone(pz_dane.zakresy_ze_skalowania("-"))

    def test_kanaly_z_uwag(self):
        self.assertEqual(pz_dane.kanaly_z_uwag("CH 001\nUz = (24,0 ± 0,1) V DC. "),
                         ("", "", ["001"]))
        self.assertEqual(pz_dane.kanaly_z_uwag("ch temp 105\nch humid 115"),
                         ("105", "115", []))


class TestPZ(unittest.TestCase):

    TEKST = (
        "Numer zlecenia laboratorium: 552/LA/T/2026;, 230/LA/TH/2026\n"
        "Obiekt wzorcowania: \n"
        "1) Przetwornik temperatury typ: KAT-W, nr fabr.: 040511513 z czujnikiem temperatury "
        "typ: Pt100, nr fabr.: 472139, \nwytwórca: RÖSSEL MESSTECHNIK. \n"
        "2) Przetwornik wilgotności względnej typ: #FPC3/5, nr fabr.: 87441, wytwórca: Galltec+mela. \n"
        "3) Przetwornik wilgotności względnej typ: #FPC3/5, nr:263051, wytwórca: Galltec+mela. \n"
        "UWAGA: \n"
        "1) Odczyt pętli prądowej 4..20 mA. \n"
        "Wzorcowanie zostanie wykonane przy napięciu zasilającym Uz = (18,0 ± 0,1) VDC. \n"
        "2) ÷ 3) Wzorcowanie zostanie wykonane przy napięciu zasilającym Uz = (24,0 ± 0,1) V DC. \n"
        " \nMetoda wzorcowania:  \n"
        "1) Metoda porównawcza w termostatach cieczowych, ILAJ 5.4/2. \n"
        "2) ÷ 3) Metoda porównawcza w komorze klimatycznej, ILAJ 5.4/11 – Wzorcowanie termohigrometrów. \n"
    )

    def setUp(self):
        self.p = pz_dane.parsuj_tekst(self.TEKST)

    def test_uwaga_nie_tworzy_przyrzadow(self):
        self.assertEqual([p.nr_fabr for p in self.p], ["040511513", "87441", "263051"])

    def test_zasilanie_w_uwagach(self):
        self.assertEqual([p.uwagi for p in self.p],
                         ["Uz = (18,0 ± 0,1) VDC.", "Uz = (24,0 ± 0,1) V DC.",
                          "Uz = (24,0 ± 0,1) V DC."])

    def test_komora(self):
        self.assertEqual([p.komora for p in self.p], [False, True, True])

    def test_skrot_lnk_do_pdf(self):
        """pliki_pz rozwija skroty .lnk do PDF-ow (zlecenie 230)."""
        from wspolne import nowa_piaskownica
        folder = nowa_piaskownica("pz_lnk")
        pdf = os.path.join(folder, "cel", "PZ 230.pdf")
        os.makedirs(os.path.dirname(pdf), exist_ok=True)
        with open(pdf, "wb") as f:
            f.write(b"%PDF-1.4\n")
        pz = os.path.join(folder, "PZ")
        os.makedirs(pz, exist_ok=True)
        try:
            import win32com.client
            sk = win32com.client.Dispatch("WScript.Shell").CreateShortcut(
                os.path.join(pz, "PZ 230 — skrót.lnk"))
            sk.TargetPath = pdf
            sk.Save()
        except Exception as e:
            self.skipTest(f"nie da sie utworzyc skrotu: {e}")
        self.assertEqual([os.path.normcase(x) for x in pz_dane.pliki_pz(pz)],
                         [os.path.normcase(pdf)])


class TestParametryKopii(unittest.TestCase):
    """Krok 3 — komorki zakladki roboczej dla kopii-przetwornika."""

    REKORD_230 = {"E": "263051", "JEDN_T": "-", "JEDN_RH": "mA",
                  "M": "-", "N": "(4 ÷ 20) mA\nodpowiada:\n(0 ÷ 100) %",
                  "P": "CH 001\nUz = (24,0 ± 0,1) V DC. "}

    def test_termohigrometr_bez_zmian(self):
        rek = {"E": "D0540007", "JEDN_T": "°C", "JEDN_RH": "%"}
        tab = pz_dane.przetworniki_z_tabeli(TABELA)
        self.assertIsNone(A.parametry_przetwornika(rek, tab, None))

    def test_230_z_protokolu(self):
        p = A.parametry_przetwornika(self.REKORD_230, {}, "-")
        self.assertEqual({k: p[k] for k in A.KOMORKI_PRZETWORNIKA if k in p},
                         {"K10": "Przetwornik", "L10": "Prąd", "K14": "1586A",
                          "K15": "-", "K16": "001", "M22": 4.0, "O22": 20.0,
                          "L28": 0.0, "O28": 100.0})
        self.assertTrue(any("multimetr" in o for o in p["_ostrzezenia"]))

    def test_230_z_tabeli(self):
        p = A.parametry_przetwornika(self.REKORD_230, pz_dane.przetworniki_z_tabeli(TABELA), "-")
        self.assertEqual((p["K14"], p["K16"], p["_zrodlo"]), ("1586A", "001", "tabela panelu"))
        self.assertEqual(p["_ostrzezenia"], [])

    def test_multimetr_ze_strony_1(self):
        p = A.parametry_przetwornika(self.REKORD_230, {}, "1586A-01")
        self.assertEqual(p["K14"], "1586A-01")

    def test_219_napiecie_z_tabeli(self):
        rek = {"E": "D0540007", "JEDN_T": "V", "JEDN_RH": "V", "P": "ch temp 105\nch humid 115"}
        p = A.parametry_przetwornika(rek, pz_dane.przetworniki_z_tabeli(TABELA), None)
        self.assertEqual((p["L10"], p["K15"], p["K16"], p["M22"], p["O22"], p["L27"], p["O27"]),
                         ("Napięcie", 105, 115, 0.0, 5.0, -40.0, 60.0))

    def test_jednostka_rozstrzyga_o_sygnale(self):
        rek = dict(self.REKORD_230, JEDN_RH="V")
        p = A.parametry_przetwornika(rek, pz_dane.przetworniki_z_tabeli(TABELA), None)
        self.assertEqual(p["L10"], "Napięcie")
        self.assertTrue(any("jednostka" in o for o in p["_ostrzezenia"]))

    def test_kanal_tekstem_z_zerami(self):
        self.assertEqual(A._kanal_do_komorki("001"), "001")
        self.assertEqual(A._kanal_do_komorki("221"), 221)
        self.assertEqual(A._kanal_do_komorki(""), "-")


class TestKrok2(unittest.TestCase):

    URZ = {"t": ("105", 30, "U"), "rh": ("115", 31, "U")}

    def test_jednostki_wiersza_19(self):
        self.assertEqual(G.jednostka_przetwornika(self.URZ, "t"), "V")
        self.assertEqual(G.jednostka_przetwornika({"t": None, "rh": ("001", 1, "I")}, "rh"), "mA")
        self.assertEqual(G.jednostka_przetwornika({"t": None, "rh": ("001", 1, "I")}, "t"), "-")

    def test_strona2_skalowanie_z_tabeli_i_uwagi_z_pz(self):
        wb = openpyxl.Workbook()
        self.addCleanup(wb.close)
        ws = wb.active
        w = G.STRONA2_PIERWSZY_WIERSZ
        ws.cell(row=w, column=16).value = "Uz = (24,0 ± 0,1) V DC."
        with redirect_stdout(io.StringIO()):
            G._wypelnij_przetworniki_strona2(
                ws, [("D0540007", [], [], {"przetwornik": self.URZ})],
                pz_dane.przetworniki_z_tabeli(TABELA))
        self.assertEqual(ws.cell(row=w, column=13).value,
                         "(0 ÷ 5) V\nodpowiada:\n(-40 ÷ 60) °C")
        self.assertEqual(ws.cell(row=w, column=14).value,
                         "(0 ÷ 5) V\nodpowiada:\n(0 ÷ 100) %")
        self.assertEqual(ws.cell(row=w, column=16).value,
                         "ch temp 105\nch humid 115\nUz = (24,0 ± 0,1) V DC.")

    def test_strona1_multimetr_analogowy(self):
        wb = openpyxl.Workbook()
        self.addCleanup(wb.close)
        ws = wb.active
        with redirect_stdout(io.StringIO()):
            G._wypelnij_strone1_przetworniki(ws, [self.URZ], "1586A-02", {}, "D0540007")
        self.assertEqual((ws["D53"].value, ws["I53"].value, ws["I58"].value),
                         ("napięcie wyjściowe przetwornika", "1586A-02", 105))

    def test_multimetr_z_naglowka(self):
        self.assertEqual(G.multimetr_z_naglowka(
            ["Multimetr wzorcowy Pt100: 1586A-02", "Data Czas;Tzadana"]), "1586A-02")


if __name__ == "__main__":
    unittest.main(verbosity=2)
