# -*- coding: utf-8 -*-
"""
Testy zaznaczania wybranych okien w plikach wynikow kroku 1.

Krok 2 wybiera dla kazdego punktu 5 wierszy reprezentacyjnych i podswietla je w
arkuszu obserwacji. Te same chwile czasowe maja byc widoczne takze w plikach
'wyniki/<serial>_wynik.xlsx' — zeby dalo sie sprawdzic, z ktorych probek loggera
powstal punkt protokolu.

Trzy rzeczy, ktorych wczesniej brakowalo:

  1. Kazdy punkt byl tu ZIELONY, takze ten, ktory w obserwacji jest POMARANCZOWY
     (kryteria stabilnosci niespelnione). Plik wynikow pokazywal wiec jako dobre
     punkty, o ktorych obserwacja ostrzegala.
  2. Nie bylo numeru punktu. W zestawieniu zbiorczym numer jest od dawna — po nim
     skacze sie po pliku w Excelu zamiast szukac kolorow wzrokiem.
  3. Oznaczenia powstawaly przy okazji wpisywania danych do protokolu, wiec plik,
     ktory pasowal czasowo, ale nie zmiescil sie w kolumnach przyrzadow protokolu
     (limit MAX_PRZYRZADY), zostawal bez zaznaczenia.
"""

import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import generuj_obserwacje as G
from wspolne import nowa_piaskownica

ZIELONY = G.FILL_DARK.fgColor.rgb
POMARANCZOWY = G.FILL_WARN_DARK.fgColor.rgb


def kolor(komorka):
    rgb = getattr(komorka.fill.fgColor, "rgb", None) if komorka.fill else None
    return None if rgb in (None, "00000000") else rgb


class TestOznaczaniePliku(unittest.TestCase):
    """Plik wynikow: kolor wg stanu punktu + numer punktu w kolumnie obok."""

    @classmethod
    def setUpClass(cls):
        cls.folder = nowa_piaskownica("oznaczenia_wynikow")

    def plik(self, nazwa, wierszy=30):
        sciezka = os.path.join(self.folder, nazwa)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Czas", "Temperatura [°C]", "Wilgotność [%RH]"])
        for i in range(wierszy):
            ws.append([f"2026-09-11 12:{i:02d}:00", 25.0 + i * 0.01, 40.0 + i * 0.1])
        wb.save(sciezka)
        return sciezka

    def oznacz(self, nazwa, wg_punktu, powody=None):
        sciezka = self.plik(nazwa)
        G._oznacz_wyniki_xlsx(sciezka, wg_punktu, powody)
        wb = openpyxl.load_workbook(sciezka)
        self.addCleanup(wb.close)
        return wb.active

    def test_punkt_bez_ostrzezenia_jest_zielony(self):
        ws = self.oznacz("zielony.xlsx", {0: {5, 6, 7, 8, 9}})
        for r in (5, 6, 7, 8, 9):
            with self.subTest(wiersz=r):
                self.assertEqual(kolor(ws.cell(row=r, column=1)), ZIELONY)

    def test_punkt_z_ostrzezeniem_jest_pomaranczowy(self):
        """W obserwacji ten punkt jest pomaranczowy — tu tez musi byc."""
        ws = self.oznacz("pomaranczowy.xlsx", {0: {5, 6, 7}},
                         powody={0: "rozrzut za duzy"})
        self.assertEqual(kolor(ws.cell(row=5, column=1)), POMARANCZOWY)

    def test_kolory_nie_mieszaja_sie_miedzy_punktami(self):
        ws = self.oznacz("mieszane.xlsx", {0: {5, 6}, 1: {10, 11}},
                         powody={1: "na styku punktow"})
        self.assertEqual(kolor(ws.cell(row=5, column=1)), ZIELONY)
        self.assertEqual(kolor(ws.cell(row=10, column=1)), POMARANCZOWY)

    def test_wiersze_poza_punktami_zostaja_biale(self):
        ws = self.oznacz("bez_reszty.xlsx", {0: {5, 6}})
        self.assertIsNone(kolor(ws.cell(row=4, column=1)))
        self.assertIsNone(kolor(ws.cell(row=7, column=1)))

    def test_kolorowany_jest_caly_wiersz(self):
        ws = self.oznacz("caly_wiersz.xlsx", {0: {5}})
        for kol in (1, 2, 3):
            with self.subTest(kolumna=kol):
                self.assertEqual(kolor(ws.cell(row=5, column=kol)), ZIELONY)

    def test_numer_punktu_w_kolumnie_znacznika(self):
        ws = self.oznacz("numery.xlsx", {0: {5, 6}, 2: {10, 11}})
        znaczniki = {ws.cell(row=r, column=5).value for r in (5, 10)}
        self.assertEqual(znaczniki, {1, 3})      # punkty liczone od 1

    def test_numer_stoi_przy_pierwszym_wierszu_punktu(self):
        ws = self.oznacz("pierwszy.xlsx", {0: {7, 8, 9}})
        self.assertEqual(ws.cell(row=7, column=5).value, 1)
        self.assertIsNone(ws.cell(row=8, column=5).value)

    def test_kolumna_znacznika_ma_naglowek(self):
        ws = self.oznacz("naglowek.xlsx", {0: {5}})
        self.assertEqual(ws.cell(row=1, column=5).value, "Nr punktu")

    def test_znacznik_nie_nadpisuje_danych(self):
        """Miedzy danymi a znacznikiem zostaje pusta kolumna — jak w zestawieniu."""
        ws = self.oznacz("odstep.xlsx", {0: {5}})
        self.assertAlmostEqual(ws.cell(row=5, column=3).value, 40.3, places=6)
        self.assertIsNone(ws.cell(row=5, column=4).value)

    def test_oznaczone_wiersze_sa_pogrubione(self):
        ws = self.oznacz("bold.xlsx", {0: {5}})
        self.assertTrue(ws.cell(row=5, column=1).font.bold)

    def test_brak_punktow_nie_rusza_pliku(self):
        sciezka = self.plik("nietkniety.xlsx")
        przed = os.path.getmtime(sciezka)
        G._oznacz_wyniki_xlsx(sciezka, {})
        self.assertEqual(os.path.getmtime(sciezka), przed)

    def test_brak_pliku_nie_wywala(self):
        G._oznacz_wyniki_xlsx(os.path.join(self.folder, "nie_ma.xlsx"), {0: {5}})


class TestZgodnoscZObserwacja(unittest.TestCase):
    """Ten sam zestaw kolorow, co w arkuszu obserwacji i w zestawieniu."""

    def test_zielony_to_ten_sam_odcien_co_w_obserwacji(self):
        self.assertEqual(G.FILL_DARK.fgColor.rgb, "00A9D08E")

    def test_pomaranczowy_to_ten_sam_odcien_co_w_obserwacji(self):
        self.assertEqual(G.FILL_WARN_DARK.fgColor.rgb, "00F4B183")

    def test_odcienie_sa_rozne(self):
        self.assertNotEqual(G.FILL_DARK.fgColor.rgb, G.FILL_WARN_DARK.fgColor.rgb)


if __name__ == "__main__":
    unittest.main(verbosity=2)
