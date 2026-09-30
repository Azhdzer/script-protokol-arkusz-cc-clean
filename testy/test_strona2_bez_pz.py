# -*- coding: utf-8 -*-
"""
Testy wypelniania Strony 2 protokolu, gdy przyrzadu NIE MA w PZ.

Zgloszenie: zlecenie wewnetrzne — PZ do niego w ogole nie powstaje. Kolumny
pomiarowe Strony 3 byly wypelnione danymi przyrzadu, a jego wiersz w tabeli na
Stronie 2 zostawal calkiem pusty, mimo ze numer fabryczny jest znany (bierze sie
z nazwy pliku wyniku, np. '10894159_wynik.xlsx').

Teraz taki wiersz dostaje to, co da sie ustalic z samego pomiaru:
    E  — nr fabryczny
    K  — rozdzielczosc odczytu temperatury (z wahania cyfr po przecinku)
    L  — rozdzielczosc odczytu wilgotnosci
    O  — nr zlecenia

Producent i typ zostaja PUSTE — z pomiaru ich nie widac, a wpisanie '-' udawaloby,
ze pole jest juz wypelnione.
"""

import os
import sys
import unittest
from contextlib import redirect_stdout
import io

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import generuj_obserwacje as G
import pz_dane

KOL = {"wytworca": 2, "typ": 4, "fabr": 5, "ewid": 6,
       "rozdz_t": 11, "rozdz_rh": 12, "zlecenie": 15}


def temperatury(krok=0.001, ile=40):
    return [round(23.0 + i * krok, 3) for i in range(ile)]


def wilgotnosci(krok=0.001, ile=40):
    return [round(45.0 + i * krok, 3) for i in range(ile)]


class TestPrzyrzadSpozaPZ(unittest.TestCase):
    """Brak PZ nie moze oznaczac pustego wiersza."""

    def wypelnij(self, uzyte, pz_mapa=None, nr_zlecenia="214"):
        wb = openpyxl.Workbook()
        ws = wb.active
        self.addCleanup(wb.close)
        with redirect_stdout(io.StringIO()) as buf:
            G.wypelnij_strone2_z_pz(ws, uzyte, pz_mapa or {}, [],
                                    nr_zlecenia=nr_zlecenia)
        self.log = buf.getvalue()
        return ws

    def komorka(self, ws, wiersz, nazwa):
        return ws.cell(row=wiersz, column=KOL[nazwa]).value

    def setUp(self):
        self.w = G.STRONA2_PIERWSZY_WIERSZ

    def test_numer_fabryczny_wpisany(self):
        ws = self.wypelnij([("10894159", temperatury(), wilgotnosci())])
        self.assertEqual(self.komorka(ws, self.w, "fabr"), "10894159")

    def test_nr_zlecenia_wpisany(self):
        ws = self.wypelnij([("10894159", temperatury(), wilgotnosci())])
        self.assertEqual(self.komorka(ws, self.w, "zlecenie"), "214")

    def test_rozdzielczosc_liczona_z_danych(self):
        ws = self.wypelnij([("10894159", temperatury(0.1), wilgotnosci(0.1))])
        self.assertEqual(self.komorka(ws, self.w, "rozdz_t"), 0.1)
        self.assertEqual(self.komorka(ws, self.w, "rozdz_rh"), 0.1)

    def test_producent_i_typ_zostaja_puste(self):
        """Z pomiaru ich nie widac — '-' udawaloby wypelnione pole."""
        ws = self.wypelnij([("10894159", temperatury(), wilgotnosci())])
        self.assertIsNone(self.komorka(ws, self.w, "wytworca"))
        self.assertIsNone(self.komorka(ws, self.w, "typ"))

    def test_przyrzad_tylko_temperatura(self):
        """Bez danych RH rozdzielczosc wilgotnosci bierze sie z temperatury."""
        ws = self.wypelnij([("10894159", temperatury(0.1), [])])
        self.assertEqual(self.komorka(ws, self.w, "rozdz_t"), 0.1)
        self.assertEqual(self.komorka(ws, self.w, "rozdz_rh"), 0.1)

    def test_kazdy_przyrzad_w_swoim_wierszu(self):
        ws = self.wypelnij([("10894159", temperatury(), wilgotnosci()),
                            ("10898381", temperatury(), wilgotnosci())])
        self.assertEqual(self.komorka(ws, self.w, "fabr"), "10894159")
        self.assertEqual(self.komorka(ws, self.w + 1, "fabr"), "10898381")

    def test_bez_numeru_zlecenia_kolumna_zostaje_pusta(self):
        ws = self.wypelnij([("10894159", temperatury(), wilgotnosci())],
                           nr_zlecenia="")
        self.assertIsNone(self.komorka(ws, self.w, "zlecenie"))

    def test_log_mowi_co_zostalo_do_recznego_uzupelnienia(self):
        self.wypelnij([("10894159", temperatury(), wilgotnosci())])
        self.assertIn("uzupelnij recznie", self.log)
        self.assertIn("10894159", self.log)


class TestPrzyrzadZPZBezZmian(unittest.TestCase):
    """
    Regresja: wiersz przyrzadu ZNALEZIONEGO w PZ ma wygladac tak jak dotad —
    nowa sciezka dotyczy wylacznie przypadku 'brak dopasowania'.
    """

    def setUp(self):
        self.w = G.STRONA2_PIERWSZY_WIERSZ
        przyrzad = pz_dane.PZPrzyrzad(
            nr_zlecenia="211", wytworca="Termoprodukt", typ="TERMIOPLUS",
            nr_fabr="2950722", nr_ewid="CLDK/B-49", komora=True)
        mapa = {pz_dane.normalizuj_serial("2950722"): przyrzad}
        wb = openpyxl.Workbook()
        self.addCleanup(wb.close)
        self.ws = wb.active
        with redirect_stdout(io.StringIO()):
            G.wypelnij_strone2_z_pz(self.ws, [("2950722", temperatury(0.1),
                                               wilgotnosci(0.1))],
                                    mapa, [], nr_zlecenia="999")

    def komorka(self, nazwa):
        return self.ws.cell(row=self.w, column=KOL[nazwa]).value

    def test_dane_ida_z_pz(self):
        self.assertEqual(self.komorka("wytworca"), "Termoprodukt")
        self.assertEqual(self.komorka("typ"), "TERMIOPLUS")
        self.assertEqual(self.komorka("fabr"), "2950722")
        self.assertEqual(self.komorka("ewid"), "CLDK/B-49")

    def test_nr_zlecenia_z_pz_ma_pierwszenstwo(self):
        """Numer z PZ, a nie ten podany jako zapasowy."""
        self.assertEqual(self.komorka("zlecenie"), "211")

    def test_puste_pola_pz_daja_myslnik(self):
        """Dotychczasowe zachowanie dla wierszy z PZ: brak danych -> '-'."""
        self.assertEqual(self.ws.cell(row=self.w, column=7).value, "-")


class TestSerialZNazwyPliku(unittest.TestCase):
    """
    Nr fabryczny z nazwy pliku wynikow to KLUCZ dopasowania przyrzadu do PZ.
    Gdy klucz jest zly, wiersz na Stronie 2 zostaje niewypelniony — mimo ze i PZ,
    i dane sa w porzadku.

    Programy loggerow dokleja do nazwy date i godzine w roznych miejscach:
    LogSoft na KONCU, eksport Testo na POCZATKU. Przedrostek nie byl odcinany,
    wiec kluczem stawal sie caly '2026-09-30-07-15-08 85517778'.
    """

    def serial(self, nazwa):
        return G._serial_z_wyniku(nazwa)

    def test_data_na_poczatku_jest_odcinana(self):
        self.assertEqual(
            self.serial("2026-09-30-07-15-08 85517778_wynik.xlsx"), "85517778")

    def test_data_na_koncu_nadal_odcinana(self):
        self.assertEqual(
            self.serial("1970325 2026-07-31 12.19.00_wynik.xlsx"), "1970325")

    def test_sama_nazwa_bez_daty_bez_zmian(self):
        self.assertEqual(self.serial("TMM230200349_wynik.xlsx"), "TMM230200349")

    def test_numer_kolejnego_wzorcowania_odcinany(self):
        self.assertEqual(self.serial("37025105_2_wynik.xlsx"), "37025105")

    def test_data_bez_godziny_na_poczatku(self):
        self.assertEqual(self.serial("2026-09-30 85517778_wynik.xlsx"), "85517778")

    def test_serial_z_literami_po_dacie(self):
        self.assertEqual(
            self.serial("2026-09-30-07-15-08 TMM2302_wynik.xlsx"), "TMM2302")

    def test_sama_data_zostaje_jak_byla(self):
        """Nie wolno zjesc calej nazwy — pusty klucz nie dopasuje sie do niczego."""
        self.assertEqual(self.serial("2026-09-30_wynik.xlsx"), "2026-09-30")


if __name__ == "__main__":
    unittest.main(verbosity=2)
