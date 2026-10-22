# -*- coding: utf-8 -*-
"""
Przetworniki w pliku CC-04 (zlecenie 219).

Naglowek pliku:
    Czujnik wzorcowy: Pt100-03; Wejscie pomiarowe Ch: 101
    Czujnik wzorcowy: Pt100-07; Wejscie pomiarowe Ch: 103
    Przetwornik U; Wejscie pomiarowe Ch: 105
    Przetwornik U; Wejscie pomiarowe Ch: 115

PTU303 rejestruje swoje odczyty (plik logera) i ma wyjscia 0-5 V mierzone
multimetrem: Ch105 = temperatura, Ch115 = wilgotnosc. To jeden przyrzad
wzorcowany na dwa sposoby. Skrypt:
  * sadzal Ch105 na miejscu CZUJNIKA WZORCOWEGO (kolumna M Strony 3, wiersz
    38 Strony 1 jako 'CC-04-LD'),
  * gubil Ch115 — kanal spoza ukladu 101-108 — w obserwacji i w protokole.
Teraz przetworniki sa osobnymi kolumnami obserwacji i osobnym przyrzadem
w protokole (para kolumn t/RH na Stronie 3, wiersz na Stronie 2).
"""

import io
import os
import sys
import unittest
from contextlib import redirect_stdout

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import generuj_obserwacje as G
from wspolne import nowa_piaskownica

NAGLOWEK = [
    "Wersja programu: OdczytTH ver. 2.12.5",
    "Zlecenie: 219",
    "Data rozpoczęcia pomiaru: 2026-10-19",
    "Pomiary wykonuje: X",
    "Komora klimatyczna: CC-04",
    "Higrometr punktu rosy: S8000",
    "Multimetr wzorcowy Pt100: 1586A-02",
    "Czujnik wzorcowy: Pt100-03; Wejście pomiarowe Ch: 101",
    "Czujnik wzorcowy: Pt100-07; Wejście pomiarowe Ch: 103",
    "Przetwornik U; Wejście pomiarowe Ch: 105",
    "Przetwornik U; Wejście pomiarowe Ch: 115",
    "Rejestrator warunków środowiskowych: MX1101-02",
    "",
    "Data Czas;Tzadana;RHzadana;Todczytana;RHodczytana;Ch101;Ch103;Ch105;Ch115;tdp;thigro;"
    "%rh;tempCh101;tempCh103;tempCh105;tempCh115;roztdp(15min);roztempCh101;roztempCh103;"
    "roztempCh105;roztempCh115;",
]


def wiersz(t, rh, i):
    u_t = (t + 40) / 100 * 5          # 0-5 V = -40..60 °C
    u_rh = rh / 100 * 5               # 0-5 V = 0..100 %
    return (f"2026-10-19 14:{i:02d}:00;{t};{rh};{t};{rh};103.85;103.88;{u_t:.4f};{u_rh:.4f};"
            f"-9.8;{t};{rh};{t};{t};brak;brak;0.05;0.02;0.02;brak;brak")


PRZEBIEG = [(10, 25), (10, 45), (20, 15), (20, 95), (40, 45), (60, 15)]


def plik(folder, linie=None):
    sciezka = os.path.join(folder, "2026-10-19 14.15_219.txt")
    dane = [wiersz(t, rh, i) for i, (t, rh) in enumerate(PRZEBIEG)]
    with open(sciezka, "w", encoding="utf-8") as f:
        f.write("\n".join((linie or NAGLOWEK) + dane) + "\n")
    return sciezka


class TestNaglowek(unittest.TestCase):

    def test_kanaly_przetwornikow(self):
        self.assertEqual(G.kanaly_przetwornikow(NAGLOWEK), [("105", "U"), ("115", "U")])

    def test_przetwornik_pradowy(self):
        linie = ["Przetwornik I; Wejście pomiarowe Ch: 107"] + NAGLOWEK[-1:]
        self.assertEqual(G.kanaly_przetwornikow(linie), [("107", "I")])

    def test_bez_przetwornikow(self):
        bez = [l for l in NAGLOWEK if not l.startswith("Przetwornik")]
        self.assertEqual(G.kanaly_przetwornikow(bez), [])

    def test_przetwornik_nie_zajmuje_slotu_czujnika(self):
        kanaly = G.kanaly_z_naglowka(NAGLOWEK)
        mapa = G.przypisz_kanaly_pliku(kanaly, G.kanaly_przetwornikow(NAGLOWEK))
        self.assertEqual(mapa, {101: "101", 103: "103"})

    def test_bez_przetwornikow_uklad_bez_zmian(self):
        kanaly = G.kanaly_z_naglowka(NAGLOWEK)
        self.assertEqual(G.przypisz_kanaly_pliku(kanaly),
                         {101: "101", 103: "103", 105: "105"})


class TestParsowanie(unittest.TestCase):

    def setUp(self):
        self.folder = nowa_piaskownica("przetworniki_cc04")
        self.nazwy, self.rows = G.parse_txt_cc04(plik(self.folder))
        self.idx = {n: i for i, n in enumerate(G.CC04_KOLUMNY)}

    def test_czujniki_wzorcowe_tylko_pt100(self):
        self.assertEqual(self.nazwy, ["Pt100-03", "Pt100-07", "", ""])

    def test_slot_105_pusty(self):
        self.assertEqual(self.rows[0][self.idx["Ch105"]], "")

    def test_przetworniki_za_ukladem(self):
        n = len(G.CC04_KOLUMNY)
        self.assertEqual(self.rows[0][n], "2.5000")       # Ch105: 10 °C
        self.assertEqual(self.rows[0][n + 1], "1.2500")   # Ch115: 25 %

    def test_sloty_z_danymi_bez_przetwornika(self):
        self.assertEqual(G.sloty_z_danymi(self.rows), {101, 103})

    def test_rola_wyjsc_z_przebiegu(self):
        n = len(G.CC04_KOLUMNY)
        self.assertEqual(G.rola_przetwornika(self.rows, n), "t")
        self.assertEqual(G.rola_przetwornika(self.rows, n + 1), "rh")

    def test_rola_niezalezna_od_kolejnosci_kanalow(self):
        """Gdyby wilgotnosc byla na nizszym kanale — rozstrzyga przebieg, nie numer."""
        linie = NAGLOWEK[:9] + ["Przetwornik U; Wejście pomiarowe Ch: 115",
                                "Przetwornik U; Wejście pomiarowe Ch: 105"] + NAGLOWEK[11:]
        _n, rows = G.parse_txt_cc04(plik(self.folder, linie))
        urz = G.urzadzenia_przetwornikow(G.kanaly_przetwornikow(linie), rows)
        self.assertEqual(urz[0]["t"][0], "105")
        self.assertEqual(urz[0]["rh"][0], "115")

    def test_jeden_przyrzad_z_dwoch_wyjsc(self):
        urz = G.urzadzenia_przetwornikow(G.kanaly_przetwornikow(NAGLOWEK), self.rows)
        self.assertEqual(len(urz), 1)
        self.assertEqual(G.opis_kanalow_przetwornika(urz[0]), "ch temp 105\nch humid 115")


class TestObserwacja(unittest.TestCase):

    def test_naglowki(self):
        self.assertEqual(G._naglowki_przetwornikow([("105", "U"), ("7", "I")]),
                         ["Wskazania multimetru Przetwornik U Ch105 [V]",
                          "Wskazania multimetru Przetwornik I Ch7 [mA]"])

    def test_kolumny_zaraz_za_wskazaniami_multimetru(self):
        n = len(G.CC04_KOLUMNY)
        i = {k: v for v, k in enumerate(G.CC04_KOLUMNY)}
        kol = [0, i["Ch101"], i["Ch103"], i["tdp"], i["tempCh101"], n, n + 1]
        wynik = G._kolejnosc_z_przetwornikami(kol)
        self.assertEqual(wynik, [0, i["Ch101"], i["Ch103"], n, n + 1, i["tdp"], i["tempCh101"]])


class TestProtokol(unittest.TestCase):

    def setUp(self):
        folder = nowa_piaskownica("przetworniki_cc04_prot")
        _n, self.rows = G.parse_txt_cc04(plik(folder))
        self.urz = G.urzadzenia_przetwornikow(G.kanaly_przetwornikow(NAGLOWEK), self.rows)

    def test_strona3_napiecia_w_parze_kolumn(self):
        wb = openpyxl.Workbook()
        self.addCleanup(wb.close)
        ws = wb.active
        rep = [([0, 1, 2, 3, 4], None)]
        G._wpisz_przetwornik(ws, rep, self.rows, self.urz[0], 21, 20, 5)
        self.assertEqual(ws["U20"].value, 2.5)
        self.assertEqual(ws["V20"].value, 1.25)
        self.assertEqual(ws["V24"].value, 2.25)      # 40 °C / 45 %
        self.assertEqual(ws["U20"].number_format, "0.0000")

    def test_strona2_wiersz_przetwornika(self):
        wb = openpyxl.Workbook()
        self.addCleanup(wb.close)
        ws = wb.active
        w = G.STRONA2_PIERWSZY_WIERSZ + 1
        for kol in (11, 12, 13, 14):
            ws.cell(row=w, column=kol).value = "-"
        uzyte = [("D0540007", [], []), ("D0540007", [], [], {"przetwornik": self.urz[0]})]
        with redirect_stdout(io.StringIO()) as buf:
            G._wypelnij_przetworniki_strona2(ws, uzyte)
        self.assertEqual([ws.cell(row=w, column=k).value for k in (11, 12, 13, 14)],
                         ["-", "-", None, None])
        self.assertEqual(ws.cell(row=w, column=16).value, "ch temp 105\nch humid 115")
        self.assertIn("UZUPELNIJ skalowanie", buf.getvalue())
        self.assertIsNone(ws.cell(row=w - 1, column=16).value)   # logger bez zmian

    def test_numer_fabryczny_z_jedynego_przyrzadu_w_komorze(self):
        import pz_dane
        pz = [pz_dane.PZPrzyrzad(nr_fabr="D0540007", komora=True),
              pz_dane.PZPrzyrzad(nr_fabr="D0540007", komora=False)]
        self.assertEqual(G._serial_przetwornika(pz), "D0540007")

    def test_kilka_przyrzadow_w_komorze_nie_zgaduje(self):
        import pz_dane
        pz = [pz_dane.PZPrzyrzad(nr_fabr="A1", komora=True),
              pz_dane.PZPrzyrzad(nr_fabr="B2", komora=True)]
        self.assertEqual(G._serial_przetwornika(pz), "")


if __name__ == "__main__":
    unittest.main(verbosity=2)
