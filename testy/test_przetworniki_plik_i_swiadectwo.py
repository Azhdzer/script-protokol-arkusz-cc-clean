# -*- coding: utf-8 -*-
"""
Krok 2a (przetworniki z pliku multimetru) i swiadectwo przetwornika — zlecenie 230.

Przetworniki Galltec #FPC3/5 (4-20 mA) mierzone 1586A na kanalach 001 i 221;
pomiar w osobnym pliku 'dat00001.csv':

    Record #;Time;Ch 001 (ADC);Ch 221 (ADC);Alarm Output;Signature
    1;21.10.2026 07:53;-0.00000002;0.00669723;0;GRAIW

ADC = prad w AMPERACH. Operator przepisywal recznie odczyty w mA do protokolu,
biorac probke z NAJBLIZSZEJ minuty (13:21:59 -> 13:22). Teraz robi to krok 2,
a krok 3 wystawia swiadectwo przetwornika (tylko RH) — porownane z wydanym
1129/230/LA/TH/2026.
"""

import datetime
import io
import os
import sys
import unittest
from contextlib import redirect_stdout

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import generuj_arkusze as A
import generuj_obserwacje as G
import pz_dane
from wspolne import KORZEN, nowa_piaskownica

CSV = ("Record #;Time;Ch 001 (ADC);Ch 221 (ADC);Alarm Output;Signature;;\n"
       "1;21.10.2026 13:20;0.00837127;0.00846282;0;GRAIW;;\n"
       "2;21.10.2026 13:21;0.00837241;0.00846128;0;VcDRK;;23.10.2026\n"
       "3;21.10.2026 13:22;0.00836627;0.00845676;0;VAYXJ;;\n")


def plik_csv(tresc=CSV):
    folder = nowa_piaskownica("przetworniki_plik")
    p = os.path.join(folder, "dat00001.csv")
    with open(p, "w", encoding="utf-8") as f:
        f.write(tresc)
    return p


def wiersz(czas):
    return [czas, "23", "28", "23.0", "28.1", "109.16", "4.6", "23.2", "28", "23.27", "0.1", "0.02"]


class TestPlikMultimetru(unittest.TestCase):

    def test_ampery_na_miliampery(self):
        czasy, kanaly = G.wczytaj_plik_multimetru(plik_csv())
        self.assertEqual(len(czasy), 3)
        self.assertEqual(set(kanaly), {"001", "221"})
        wart, typ = kanaly["001"]
        self.assertEqual(typ, "I")
        self.assertAlmostEqual(wart[0], 8.37127, places=6)

    def test_kolumna_mA_ma_pierwszenstwo(self):
        tresc = ("Record #;Time;Ch 001 (ADC);Ch 001 (mA)\n"
                 "1;21.10.2026 13:20;0.00837;8.371267\n")
        _c, kanaly = G.wczytaj_plik_multimetru(plik_csv(tresc))
        self.assertAlmostEqual(kanaly["001"][0][0], 8.371267, places=6)

    def test_napiecie(self):
        tresc = "Record #;Time;Ch 105 (VDC)\n1;21.10.2026 13:20;2.4952\n"
        _c, kanaly = G.wczytaj_plik_multimetru(plik_csv(tresc))
        self.assertEqual(kanaly["105"][1], "U")

    def test_najblizsza_minuta(self):
        """13:20:01 -> 13:20, 13:21:59 -> 13:22 (tak dobieral operator)."""
        rows = [wiersz("2026-10-21 13:20:01"), wiersz("2026-10-21 13:21:59"),
                wiersz("2026-10-21 14:00:00")]
        with redirect_stdout(io.StringIO()):
            kan = G.dolacz_przetworniki_z_plikow(rows, [plik_csv()])
        self.assertEqual([(k, t, i) for k, t, i, _p in kan], [("001", "I", 12), ("221", "I", 13)])
        self.assertEqual([r[12] for r in rows], ["8.371270", "8.366270", ""])

    def test_naglowki_obserwacji(self):
        self.assertEqual(G._naglowki_przetwornikow([("001", "I", 12, "dat00001.csv")]),
                         ["Wskazania multimetru Przetwornik I Ch001 [mA] (dat00001.csv)"])

    def test_kolumny_za_ch001_w_cc(self):
        self.assertEqual(G._kolejnosc_z_przetwornikami([0, 5, 6, 12, 13], n=12, za_kanalami=5.5),
                         [0, 5, 12, 13, 6])


class TestPrzyrzadyZTabeli(unittest.TestCase):

    TABELA = pz_dane.przetworniki_z_tabeli([
        ["263051", "-", "001", "1586A", "Prąd", "4", "20", "", "", "0", "100"],
        ["87441", "-", "221", "1586A", "Prąd", "4", "20", "", "", "0", "100"]])

    def test_kanal_do_przyrzadu_wg_tabeli(self):
        kan = [("001", "I", 12, "dat.csv"), ("221", "I", 13, "dat.csv")]
        urz = G.urzadzenia_przetwornikow(kan, [], self.TABELA)
        self.assertEqual([(u["serial"], u["t"], u["rh"][0]) for u in urz],
                         [("263051", None, "001"), ("87441", None, "221")])

    def test_opis_jednego_kanalu(self):
        self.assertEqual(G.opis_kanalow_przetwornika({"t": None, "rh": ("001", 12, "I")}),
                         "CH 001")


class TestSwiadectwo(unittest.TestCase):

    W253 = {"D": 23.59, "E": 28.9, "F": "-", "G": 8.365, "H": "-", "I": 27.3,
            "J": "-", "K": 1.5999999999999979, "M": "-", "O": 1.0}

    def test_formaty_jak_w_wydanym(self):
        p = A.kalibracja_przetwornika(self.W253, "23, 30")
        self.assertEqual((p["wartosc_odn"], p["wartosc_odn_RH"], p["prad_RH"],
                          p["zmierzona_RH"], p["poprawka_RH"], p["niepewnosc_RH"]),
                         ("23.59", "28.9", "8.365", "27.3", "1.6", "1.0"))
        self.assertEqual(p["zmierzona"], "-")

    def test_placeholdery_przetwornika(self):
        rek = {"F": "-", "P": "CH 001\nUz = (24,0 ± 0,1) V DC. ",
               "PRZETWORNIK": {"L10": "Prąd", "M22": 4.0, "O22": 20.0, "L28": 0.0, "O28": 100.0}}
        ph = A.placeholdery_przetwornika(rek)
        self.assertEqual((ph["[symbol_wy]"], ph["[jedn_wy]"], ph["[wyjscie_dopelniacz]"],
                          ph["[wyjscie_przym]"], ph["[wy_min]"], ph["[wy_max]"],
                          ph["[rh_min]"], ph["[rh_max]"], ph["[zasilanie]"], ph["[nr_ewid_tekst]"]),
                         ("I", "mA", "prądu", "prądowe", "4", "20", "0", "100",
                          "Uz = (24,0 ± 0,1) V DC.", ""))

    def test_napiecie(self):
        ph = A.placeholdery_przetwornika({"PRZETWORNIK": {"L10": "Napięcie"}})
        self.assertEqual((ph["[symbol_wy]"], ph["[jedn_wy]"], ph["[wyjscie_dopelniacz]"]),
                         ("U", "V", "napięcia"))

    def test_szablon_ma_placeholdery_i_brak_danych_klienta(self):
        sciezka = os.path.join(KORZEN, "xxx_yyy_LA_TH_2026 - przetwornik RH.docx")
        if not os.path.exists(sciezka):
            self.skipTest("brak szablonu swiadectwa przetwornika w folderze roboczym")
        import docx
        d = docx.Document(sciezka)
        tekst = "\n".join(p.text for p in d.paragraphs)
        tabele = "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
        for ph in ("[zleceniodawca]", "[typ]", "[nr_fabr]", "[wytworca]", "[prad_1_RH]",
                   "[zmierzona_1_RH]", "[symbol_wy]", "[jedn_wy]", "[zasilanie]", "[histereza]"):
            with self.subTest(ph=ph):
                self.assertIn(ph, tekst + tabele)
        self.assertNotIn("263051", tekst + tabele)


class TestHistereza(unittest.TestCase):
    """Nota o histerezie: punkt z PZ (nazwa zakladki), nie zaokraglony odczyt."""

    def doc(self):
        import docx
        d = docx.Document()
        d.add_paragraph("[histereza]")
        return d

    def test_punkt_z_pz(self):
        A.ZAKLADKI_Z_PZ.add("23, 60")
        self.addCleanup(A.ZAKLADKI_Z_PZ.discard, "23, 60")
        d = self.doc()
        A._zastap_histereza_w_dok(d, [{"nazwa": "23, 60 (2)", "wartosc_odn": "23.83",
                                       "wartosc_odn_RH": "57.6"}])
        self.assertIn("Punkt (23 °C, 60 %rh)", d.paragraphs[0].text)

    def test_bez_pz_jak_dotad(self):
        d = self.doc()
        A._zastap_histereza_w_dok(d, [{"nazwa": "20, 44 (2)", "wartosc_odn": "19.9",
                                       "wartosc_odn_RH": "44.6"}])
        self.assertIn("Punkt (20 °C, 40 %rh)", d.paragraphs[0].text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
