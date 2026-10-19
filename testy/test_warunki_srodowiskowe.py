# -*- coding: utf-8 -*-
"""
Testy kroku 4 — warunki srodowiskowe w protokole (wypelnij_warunki.py).

Krok 3 wpisywal t/RH pomieszczenia (Strona 3, kolumny F/G) tylko dla blokow,
ktore trafily do jakiejs kopii arkusza. Punkt wyszarzony dla wszystkich
przyrzadow — np. spoza ich zamowienia — do zadnej kopii nie trafial, wiec jego
warunki zostawaly puste (zlecenie 201: punkt 35 °C / 60 %).

Krok 4 pracuje wprost na protokole: wypelnia KAZDY punkt, takze ten, w ktorym
zaden przyrzad nie mial odczytow. Excela nie uruchamia.
"""

import datetime
import io
import os
import sys
import unittest
from contextlib import redirect_stdout

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import cc_config as C
import generuj_arkusze as A
import wypelnij_warunki as W
from wspolne import nowa_piaskownica

# Punkty protokolu: (data, start, koniec)
PUNKTY = [
    ("15.10.2026", "18:40", "18:44"),
    ("15.10.2026", "21:52", "21:56"),
    ("16.10.2026", "17:09", "17:13"),   # punkt bez odczytow przyrzadu — tez ma byc wypelniony
]


def protokol(sciezka, punkty=PUNKTY, wypelnione=None):
    """Protokol z blokami punktow jak w szablonie (od wiersza 20, po 5 wierszy)."""
    wb = openpyxl.Workbook()
    wb.active.title = "Strona 1"
    wb.create_sheet("Strona 2")
    ws = wb.create_sheet("Strona 3")
    for i, (data, start, koniec) in enumerate(punkty):
        r0 = A.START_ROW_S3 + i * A.BLOK_S3
        ws.cell(row=r0, column=1).value = i + 1
        ws.cell(row=r0 + 1, column=5).value = data
        ws.cell(row=r0 + 2, column=5).value = start
        ws.cell(row=r0 + 3, column=5).value = koniec
    for (blok, offset), (t, rh) in (wypelnione or {}).items():
        r = A.START_ROW_S3 + blok * A.BLOK_S3 + offset
        ws.cell(row=r, column=6).value = t
        ws.cell(row=r, column=7).value = rh
    wb.save(sciezka)
    return sciezka


def plik_czujnika(folder, przerwa=None):
    """
    Plik pomieszczenia w ukladzie oryginalu: arkusz "10'26", naglowek w 4. wierszu,
    rekord co minute. `przerwa` = (od, do) — okres bez zapisu czujnika.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "10'26"
    ws.append(["Pomieszczenie nr 9 (MX1101-02)"])
    ws.append([])
    ws.append([])
    ws.append(["Data Czas", "t °C", "RH %"])
    t = datetime.datetime(2026, 10, 15, 0, 0, 6)
    while t < datetime.datetime(2026, 10, 17):
        if not (przerwa and przerwa[0] <= t <= przerwa[1]):
            ws.append([t, 22.0 + t.hour / 100, 45.0 + t.minute / 100])
        t += datetime.timedelta(minutes=1)
    wb.save(os.path.join(folder, f"Pom. nr {A.NR_POMIESZCZENIA} ({A.MODEL_CZUJNIKA}) - 2026.xlsx"))


class ZPiaskownica(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.folder = nowa_piaskownica("warunki_srodowiskowe")
        plik_czujnika(cls.folder)
        cls._stary = A.FOLDER
        A.FOLDER = cls.folder

    @classmethod
    def tearDownClass(cls):
        A.FOLDER = cls._stary

    def wypelnij(self, nazwa, **kw):
        sciezka = os.path.join(self.folder, nazwa)
        with redirect_stdout(io.StringIO()) as buf:
            wynik = W.wypelnij_protokol(sciezka, **kw)
        self.log = buf.getvalue()
        return sciezka, wynik

    def fg(self, sciezka, blok, offset):
        wb = openpyxl.load_workbook(sciezka)
        try:
            ws = wb["Strona 3"]
            r = A.START_ROW_S3 + blok * A.BLOK_S3 + offset
            return ws.cell(row=r, column=6).value, ws.cell(row=r, column=7).value
        finally:
            wb.close()


class TestWypelnianie(ZPiaskownica):

    def test_kazdy_punkt_dostaje_warunki(self):
        """Takze punkt 3 — ten, w ktorym przyrzad nie mial odczytow."""
        sciezka, w = self.wypelnij(os.path.basename(protokol(
            os.path.join(self.folder, "wszystkie.xlsx"))))
        self.assertEqual(w["wpisane"], 6)            # 3 punkty x (start + koniec)
        for blok in range(3):
            for offset in (2, 3):
                with self.subTest(blok=blok, offset=offset):
                    t, rh = self.fg(sciezka, blok, offset)
                    self.assertIsNotNone(t)
                    self.assertIsNotNone(rh)

    def test_wartosc_z_najblizszej_chwili(self):
        """Start 18:40 — rekord czujnika 18:40:06, t = 22 + 18/100."""
        sciezka, _w = self.wypelnij(os.path.basename(protokol(
            os.path.join(self.folder, "chwila.xlsx"))))
        t, rh = self.fg(sciezka, 0, 2)
        self.assertAlmostEqual(t, 22.18, places=2)
        self.assertAlmostEqual(rh, 45.40, places=2)

    def test_start_i_koniec_osobno(self):
        sciezka, _w = self.wypelnij(os.path.basename(protokol(
            os.path.join(self.folder, "start_koniec.xlsx"))))
        self.assertNotEqual(self.fg(sciezka, 0, 2)[1], self.fg(sciezka, 0, 3)[1])

    def test_wypelnione_reczne_zostaja(self):
        """Poprawione recznie wartosci nie moga zniknac przy kolejnym przebiegu."""
        sciezka, w = self.wypelnij(os.path.basename(protokol(
            os.path.join(self.folder, "reczne.xlsx"), wypelnione={(0, 2): (99.9, 11.1)})))
        self.assertEqual(self.fg(sciezka, 0, 2), (99.9, 11.1))
        self.assertEqual(w["pominiete"], 1)
        self.assertEqual(w["wpisane"], 5)

    def test_nadpisywanie_na_zadanie(self):
        sciezka, _w = self.wypelnij(os.path.basename(protokol(
            os.path.join(self.folder, "nadpisz.xlsx"), wypelnione={(0, 2): (99.9, 11.1)})),
            nadpisuj=True)
        self.assertNotEqual(self.fg(sciezka, 0, 2), (99.9, 11.1))

    def test_nic_do_zrobienia_nie_zapisuje_pliku(self):
        sciezka = protokol(os.path.join(self.folder, "pelny.xlsx"),
                           punkty=PUNKTY[:1],
                           wypelnione={(0, 2): (1.0, 2.0), (0, 3): (3.0, 4.0)})
        przed = os.path.getmtime(sciezka)
        _s, w = self.wypelnij("pelny.xlsx")
        self.assertEqual(w["wpisane"], 0)
        self.assertEqual(os.path.getmtime(sciezka), przed)


class TestBrakDanych(unittest.TestCase):
    """Gdy czujnik nie ma zapisu blisko punktu — komorka zostaje pusta."""

    @classmethod
    def setUpClass(cls):
        cls.folder = nowa_piaskownica("warunki_przerwa")
        plik_czujnika(cls.folder, przerwa=(datetime.datetime(2026, 10, 15, 18, 0),
                                           datetime.datetime(2026, 10, 15, 19, 30)))
        cls._stary = A.FOLDER
        A.FOLDER = cls.folder

    @classmethod
    def tearDownClass(cls):
        A.FOLDER = cls._stary

    def test_przerwa_w_zapisie_zostawia_puste(self):
        """Lepiej uzupelnic recznie niz wpisac warunki z innej godziny."""
        sciezka = protokol(os.path.join(self.folder, "przerwa.xlsx"))
        with redirect_stdout(io.StringIO()) as buf:
            w = W.wypelnij_protokol(sciezka)
        self.assertEqual(w["brak"], 2)               # start i koniec punktu 1
        self.assertIn("BRAK", buf.getvalue())
        wb = openpyxl.load_workbook(sciezka)
        self.assertIsNone(wb["Strona 3"].cell(row=A.START_ROW_S3 + 2, column=6).value)
        wb.close()


class TestNiepoprawneWejscie(ZPiaskownica):

    def test_plik_bez_strony_3(self):
        sciezka = os.path.join(self.folder, "nie_protokol.xlsx")
        openpyxl.Workbook().save(sciezka)
        with self.assertRaises(ValueError):
            W.wypelnij_protokol(sciezka)


class TestPanel(unittest.TestCase):
    """Krok w rejestrze i w panelu — osobny, poza 'Uruchom caly obieg'."""

    def test_krok_w_rejestrze(self):
        self.assertIn("srod", [k for k, _e, _o in C.KROKI])

    def test_wybor_protokolu_na_stronie_kroku(self):
        u = C.WG_ENV["SR_PROTOKOL"]
        self.assertEqual((u.krok, u.poziom, u.typ), ("srod", C.PODSTAWOWY, "plik"))

    def test_ustawienia_czujnika_przy_tym_kroku(self):
        for env in ("GEN_NR_POM", "GEN_MODEL_CZUJ", "GEN_TOL_CZUJ"):
            with self.subTest(env=env):
                self.assertEqual(C.WG_ENV[env].krok, "srod")

    def test_krok_3_nadal_widzi_ustawienia_czujnika(self):
        """Przeniesione w panelu — ale krok 3 dalej je czyta (przez env)."""
        self.assertTrue(hasattr(A, "NR_POMIESZCZENIA"))
        self.assertTrue(hasattr(A, "TOLERANCJA_CZUJNIKA_MIN"))

    def test_nie_wchodzi_do_calego_obiegu(self):
        import app_gui
        self.assertIn("srod", app_gui.URUCHAMIALNE)
        with open(app_gui.__file__, encoding="utf-8") as f:
            zrodlo = f.read()
        self.assertIn('self._uruchom(["analiza", "obs", "ark"])', zrodlo)


if __name__ == "__main__":
    unittest.main(verbosity=2)
