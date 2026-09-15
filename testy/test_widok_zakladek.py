# -*- coding: utf-8 -*-
"""
Testy wygladu paska zakladek w gotowych plikach .xlsx (cc_widok.py).

Zgloszenie: po otwarciu protokolu i kopii arkusza obliczeniowego na dole widac
bylo tylko jedna zakladke — zeby dostac sie do pozostalych stron, trzeba bylo
recznie przeciagnac suwak w lewo. Problem byl juz kiedys naprawiony i wrocil.

Dwie przyczyny, obie widoczne w 'xl/workbook.xml':

  1. tabRatio znikalo z protokolu. Krok 3 otwiera protokol przez Excel COM
     (wpisanie F/G do Strony 3) i przy zapisie Excel przepisuje <workbookView>
     po swojemu, gubiac ten atrybut.
  2. firstSheet=6 w kopiach. Przewijanie paska szlo przez zywe okno Excela
     (Windows(1).ScrollWorkbookTabs) i bywalo nieskuteczne — w jednym przebiegu
     dwie kopie z czterech mialy firstSheet=0, a dwie firstSheet=6, czyli pasek
     przewiniety na sam koniec, do zakladki 'Wyniki'.

Dlatego widok ustawiamy na GOTOWYM PLIKU, podmieniajac jedna czesc archiwum
ZIP — deterministycznie i po wszystkich zapisach, takze tych przez COM.
"""

import os
import re
import shutil
import sys
import unittest
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import cc_config as C
import cc_widok as W
from wspolne import nowa_piaskownica


def atrybut(sciezka, nazwa):
    with zipfile.ZipFile(sciezka) as z:
        xml = z.read("xl/workbook.xml").decode("utf-8", "replace")
    tag = re.search(r"<workbookView\b[^>]*/?>", xml)
    if not tag:
        return None
    m = re.search(r'%s="([^"]*)"' % nazwa, tag.group(0))
    return m.group(1) if m else None


class TestPrzeliczenieSzerokosci(unittest.TestCase):
    """Panel podaje ulamek (0,85), plik chce liczbe calkowita 0..1000."""

    def test_ulamek_na_promile(self):
        self.assertEqual(W._na_promile(0.85), 850)

    def test_jeden_to_caly_pasek(self):
        self.assertEqual(W._na_promile(1.0), 1000)

    def test_wartosc_juz_w_promilach_zostaje(self):
        """Gdyby ktos poprawil plik ustawien recznie i wpisal 850."""
        self.assertEqual(W._na_promile(850), 850)

    def test_zapis_procentowy_tez_dziala(self):
        self.assertEqual(W._na_promile(85), 850)

    def test_smiec_daje_wartosc_domyslna(self):
        self.assertEqual(W._na_promile(None), W._na_promile(W.TAB_RATIO_DOMYSLNY))
        self.assertEqual(W._na_promile("abc"), W._na_promile(W.TAB_RATIO_DOMYSLNY))

    def test_poza_zakresem_jest_przycinane(self):
        self.assertEqual(W._na_promile(5000), 1000)
        self.assertEqual(W._na_promile(-1), 0)


class TestPoprawkaXml(unittest.TestCase):
    """Trzy uklady <workbookView>, ktore realnie wychodza z Excela i openpyxl."""

    def test_nadpisuje_istniejace_atrybuty(self):
        xml = ('<workbook><bookViews><workbookView xWindow="-120" tabRatio="600" '
               'firstSheet="6" activeTab="6"/></bookViews><sheets/></workbook>')
        nowy = W.popraw_workbook_xml(xml, 0.85, pierwszy=0)
        self.assertIn('tabRatio="850"', nowy)
        self.assertIn('firstSheet="0"', nowy)

    def test_dokłada_brakujacy_tabRatio(self):
        """Tak zapisuje protokol Excel po edycji przez COM — bez tabRatio."""
        xml = ('<workbook><bookViews><workbookView xWindow="-120" '
               'windowWidth="29040"/></bookViews><sheets/></workbook>')
        nowy = W.popraw_workbook_xml(xml, 0.85, pierwszy=0)
        self.assertIn('tabRatio="850"', nowy)
        self.assertIn('windowWidth="29040"', nowy)

    def test_dokłada_caly_blok_gdy_go_brak(self):
        xml = '<workbook><fileVersion/><sheets><sheet name="A"/></sheets></workbook>'
        nowy = W.popraw_workbook_xml(xml, 0.85, pierwszy=0)
        self.assertIn("<bookViews>", nowy)
        self.assertIn('tabRatio="850"', nowy)

    def test_blok_stoi_przed_lista_arkuszy(self):
        """Schemat OOXML wymaga <bookViews> przed <sheets> — inaczej Excel protestuje."""
        xml = '<workbook><fileVersion/><sheets><sheet name="A"/></sheets></workbook>'
        nowy = W.popraw_workbook_xml(xml, 0.85, pierwszy=0)
        self.assertLess(nowy.index("<bookViews>"), nowy.index("<sheets>"))

    def test_aktywna_zakladka_domyslnie_bez_zmian(self):
        """Kopia ma otwierac sie na 'Wyniki' — nie wolno tego zepsuc."""
        xml = '<workbook><bookViews><workbookView activeTab="6"/></bookViews></workbook>'
        nowy = W.popraw_workbook_xml(xml, 0.85, pierwszy=0)
        self.assertIn('activeTab="6"', nowy)

    def test_aktywna_zakladka_ustawiana_na_zadanie(self):
        xml = '<workbook><bookViews><workbookView activeTab="6"/></bookViews></workbook>'
        nowy = W.popraw_workbook_xml(xml, 0.85, pierwszy=0, aktywny=0)
        self.assertIn('activeTab="0"', nowy)

    def test_nieznany_uklad_zostaje_nietkniety(self):
        """Lepiej nic nie zrobic niz uszkodzic plik."""
        xml = "<cos_zupelnie_innego/>"
        self.assertEqual(W.popraw_workbook_xml(xml, 0.85), xml)


class TestNaPrawdziwymPliku(unittest.TestCase):
    """Podmiana czesci w archiwum ZIP nie moze niczego innego ruszyc."""

    @classmethod
    def setUpClass(cls):
        cls.folder = nowa_piaskownica("widok_zakladek")

    def plik(self, nazwa, arkuszy=3):
        sciezka = os.path.join(self.folder, nazwa)
        wb = openpyxl.Workbook()
        wb.active.title = "Strona 1"
        for i in range(2, arkuszy + 1):
            wb.create_sheet(f"Strona {i}")["A1"] = f"dane {i}"
        wb.views[0].tabRatio = 600
        wb.views[0].firstSheet = arkuszy - 1      # pasek przewiniety na koniec
        wb.save(sciezka)
        return sciezka

    def test_poprawia_szerokosc_i_przewiniecie(self):
        p = self.plik("zwykly.xlsx")
        self.assertEqual(atrybut(p, "tabRatio"), "600")
        self.assertTrue(W.wymus_widok(p, 0.85, pierwszy=0))
        self.assertEqual(atrybut(p, "tabRatio"), "850")
        self.assertEqual(atrybut(p, "firstSheet"), "0")

    def test_dane_arkuszy_zostaja_nienaruszone(self):
        p = self.plik("dane.xlsx")
        W.wymus_widok(p, 0.85, pierwszy=0)
        wb = openpyxl.load_workbook(p)
        try:
            self.assertEqual(wb.sheetnames, ["Strona 1", "Strona 2", "Strona 3"])
            self.assertEqual(wb["Strona 2"]["A1"].value, "dane 2")
        finally:
            wb.close()

    def test_pozostale_czesci_archiwum_bez_zmian(self):
        """
        Kopie arkusza obliczeniowego maja zapamietane wyniki formul — swiadectwa
        Word czytaja wlasnie je. Wolno nam ruszyc tylko 'xl/workbook.xml'.
        """
        p = self.plik("czesci.xlsx")
        kopia = p + ".przed"
        shutil.copy2(p, kopia)
        W.wymus_widok(p, 0.85, pierwszy=0)
        with zipfile.ZipFile(kopia) as a, zipfile.ZipFile(p) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for nazwa in a.namelist():
                if nazwa == W.CZESC:
                    continue
                with self.subTest(czesc=nazwa):
                    self.assertEqual(a.read(nazwa), b.read(nazwa))

    def test_plik_otwiera_sie_w_openpyxl_po_poprawce(self):
        p = self.plik("otwieralny.xlsx")
        W.wymus_widok(p, 0.85, pierwszy=0)
        wb = openpyxl.load_workbook(p)
        try:
            self.assertEqual(wb.views[0].tabRatio, 850)
            self.assertEqual(wb.views[0].firstSheet, 0)
        finally:
            wb.close()

    def test_brak_pliku_nie_wywala(self):
        self.assertFalse(W.wymus_widok(os.path.join(self.folder, "nie_ma.xlsx")))

    def test_plik_ktory_nie_jest_xlsx_nie_wywala(self):
        p = os.path.join(self.folder, "tekst.xlsx")
        with open(p, "w", encoding="utf-8") as f:
            f.write("to nie jest archiwum")
        self.assertFalse(W.wymus_widok(p))

    def test_nie_zostawia_pliku_tymczasowego(self):
        p = self.plik("sprzatanie.xlsx")
        W.wymus_widok(p, 0.85, pierwszy=0)
        self.assertFalse(os.path.exists(p + ".widok.tmp"))


class TestUstawienieWPanelu(unittest.TestCase):
    """Jedno pokretlo dla wszystkich plikow, nie osobne dla kazdego kroku."""

    def test_ustawienie_jest_wspolne_dla_calego_obiegu(self):
        self.assertEqual(C.WG_ENV["GEN_TAB_RATIO"].krok, "przygotowanie")

    def test_domyslna_zgodna_z_modulem(self):
        self.assertEqual(C.WG_ENV["GEN_TAB_RATIO"].domyslna, W.TAB_RATIO_DOMYSLNY)

    def test_domyslna_jest_wieksza_niz_excelowa(self):
        """0,6 Excela to wlasnie ten sciśnięty pasek, od ktorego sie zaczelo."""
        self.assertGreater(W._na_promile(C.WG_ENV["GEN_TAB_RATIO"].domyslna), 600)


if __name__ == "__main__":
    unittest.main(verbosity=2)
