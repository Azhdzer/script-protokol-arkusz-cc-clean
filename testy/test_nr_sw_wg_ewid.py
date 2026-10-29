# -*- coding: utf-8 -*-
"""
GEN_NR_SW_WG_EWID — numery swiadectw w kolejnosci nr EWIDENCYJNEGO.

Klienci chca, zeby swiadectwo nr 1047 dostal przyrzad U90337-18, 1048 — U90337-36
itd. (zlecenie 233), a nie przyrzad o najmniejszym nr fabrycznym (kolejnosc kopii).
Domyslnie (wylaczone) numeracja bez zmian.
"""

import io
import os
import re
import sys
import unittest
from contextlib import redirect_stdout
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import generuj_arkusze as A
from wspolne import nowa_piaskownica

try:
    from docx import Document
except ImportError:          # pragma: no cover
    Document = None


class _PZ:
    def __init__(self, nr_ewid):
        self.nr_ewid, self.uzytkownik, self.zleceniodawca = nr_ewid, "", ""
        self.wytworca, self.typ = "", ""


class TestNumerySwiadectw(unittest.TestCase):

    def test_domyslnie_kolejnosc_kopii(self):
        self.assertEqual(A.numery_swiadectw(["U90337-36", "U90337-18"], 1047), [1047, 1048])

    def test_wg_ewid_naturalnie(self):
        ewid = ["U90337-36", "U90337-9", "U90337-100", "U90337-18"]
        self.assertEqual(A.numery_swiadectw(ewid, 1047, wg_ewid=True), [1049, 1047, 1050, 1048])

    def test_bez_ewid_na_koncu(self):
        self.assertEqual(A.numery_swiadectw(["", "B-2", None, "B-1"], 10, wg_ewid=True),
                         [12, 11, 13, 10])

    def test_ewid_z_pz_gdy_strona2_pusta(self):
        with mock.patch.object(A, "_pz_mapa_arkusze", return_value={"371": _PZ("U-5")}):
            self.assertEqual(A._nr_ewid_kopii({"F": None}, "233_LA_TH_2026 - testo 174T - 371.xlsx"),
                             "U-5")
            self.assertEqual(A._nr_ewid_kopii({"F": "U-7"}, "233_LA_TH_2026 - testo 174T - 371.xlsx"),
                             "U-7")


@unittest.skipIf(Document is None, "brak python-docx")
class TestSwiadectwaWgEwid(unittest.TestCase):

    def test_word_dostaje_numer_wg_ewid(self):
        folder = nowa_piaskownica("nr_sw_wg_ewid")
        doc = Document()
        doc.add_paragraph("Swiadectwo [nr_sw] dla [nr_ewid]")
        doc.save(os.path.join(folder, "W.docx"))
        # kopie w kolejnosci nr fabrycznego, nr ewidencyjne "pomieszane"
        kopie = [f"233_LA_TH_2026 - testo 174T - 3719{i}.xlsx" for i in (1, 2, 3)]
        s2 = [{"F": "U90337-41"}, {"F": "U90337-18"}, {"F": "U90337-36"}]
        with mock.patch.object(A, "_pz_mapa_arkusze", return_value={}), \
                redirect_stdout(io.StringIO()):
            nazwy = A.utworz_kopie_word(folder, "W.docx", "W.docx", "W.docx", s2, kopie,
                                        [], [[], [], []], 1047, wg_ewid=True)
        wynik = {}
        for nazwa in nazwy:
            tekst = " ".join(p.text for p in Document(os.path.join(folder, nazwa)).paragraphs)
            wynik[re.match(r"(\d+)_", nazwa).group(1)] = tekst
        self.assertEqual(wynik, {"1047": "Swiadectwo 1047 dla U90337-18",
                                 "1048": "Swiadectwo 1048 dla U90337-36",
                                 "1049": "Swiadectwo 1049 dla U90337-41"})


if __name__ == "__main__":
    unittest.main(verbosity=2)
