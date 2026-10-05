# -*- coding: utf-8 -*-
"""
Testy kanalow multimetru w plikach komory CC-04.

Zgloszenie (zlecenie 221): obserwacja nie wziela wskazan multimetru. Pomiar szedl
w komorze CC-04, ale tylko na JEDNYM czujniku, wpietym w kanal Ch001 zamiast w
zwykle wejscia Ch101..Ch108:

    Komora klimatyczna: CC-04
    Czujnik wzorcowy: Pt100-31; Wejscie pomiarowe Ch: 001
    Data Czas;Tzadana;RHzadana;Todczytana;RHodczytana;Ch001;tdp;thigro;%rh;tempCh001;...

Parser CC-04 szukal kolumn po stalych nazwach 'Ch101', 'tempCh101', ...
Zadnej nie znalazl, wiec rezystancja i temperatura czujnika przepadaly, a do
obserwacji trafialy puste kolumny.

Plik POZOSTAJE plikiem CC-04 (tak mowi naglowek i tak jest w rzeczywistosci).
Zmienia sie tylko to, ze kanaly sa brane z naglowka pliku i sadzane na sloty
ukladu — reszta skryptu pracuje na niezmienionych kolumnach.
"""

import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import generuj_obserwacje as G
from wspolne import nowa_piaskownica

METADANE_221 = [
    "Wersja programu: OdczytTH ver. 2.12.5",
    "Zlecenie: 221",
    "Data rozpoczecia pomiaru: 2026-10-03",
    "Pomiary wykonuje: X",
    "Komora klimatyczna: CC-04",
    "Higrometr punktu rosy: ",
    "Multimetr wzorcowy Pt100: 1586A-02",
    "Czujnik wzorcowy: Pt100-31; Wejscie pomiarowe Ch: 001",
    "Rejestrator warunkow srodowiskowych: MX1101-02",
    "Rotametr: LZT-M-6-04",
    "Zasilacz: RZS-1D",
    "TH kontrolny: HC2-S2",
    "",
]
NAGLOWEK_1_KANAL = ("Data Czas;Tzadana;RHzadana;Todczytana;RHodczytana;Ch001;"
                    "tdp;thigro;%rh;tempCh001;roztdp(15min);roztempCh001;")


def plik(sciezka, metadane, naglowek, wiersze):
    with open(sciezka, "w", encoding="cp1250", newline="\r\n") as f:
        f.write("\n".join(list(metadane) + [naglowek] + list(wiersze)) + "\n")
    return sciezka


class TestKanalyZNaglowka(unittest.TestCase):

    def test_zera_wiodace_zachowane(self):
        """Kolumny w pliku nazywaja sie 'tempCh001' — po tej pisowni ich szukamy."""
        self.assertEqual(G.kanaly_z_naglowka([NAGLOWEK_1_KANAL]), ["001"])

    def test_uklad_osmiokanalowy(self):
        naglowek = "Data Czas;Tzadana;" + ";".join(f"Ch{c}" for c in range(101, 109))
        self.assertEqual(G.kanaly_z_naglowka([naglowek]),
                         [str(c) for c in range(101, 109)])

    def test_kolumny_pochodne_nie_licza_sie_jako_kanaly(self):
        """'tempCh001' i 'roztempCh001' to nie osobne kanaly."""
        self.assertEqual(len(G.kanaly_z_naglowka([NAGLOWEK_1_KANAL])), 1)

    def test_brak_naglowka(self):
        self.assertEqual(G.kanaly_z_naglowka(["cos innego"]), [])


class TestPrzypisanieDoSlotow(unittest.TestCase):

    def test_uklad_standardowy_bez_zmian(self):
        kanaly = [str(c) for c in range(101, 109)]
        mapa = G.przypisz_kanaly_pliku(kanaly)
        self.assertEqual(mapa, {c: str(c) for c in range(101, 109)})

    def test_jeden_czujnik_na_kanale_001(self):
        """Zgloszenie 221 — czujnik laduje na pierwszym slocie glownym."""
        self.assertEqual(G.przypisz_kanaly_pliku(["001"]), {101: "001"})

    def test_kilka_nietypowych_kanalow_na_kolejne_sloty_glowne(self):
        mapa = G.przypisz_kanaly_pliku(["001", "002"])
        self.assertEqual(mapa, {101: "001", 103: "002"})

    def test_czesc_standardowych_kanalow(self):
        """Podlaczone tylko 105 i 107 — zostaja na swoich miejscach."""
        self.assertEqual(G.przypisz_kanaly_pliku(["105", "107"]),
                         {105: "105", 107: "107"})

    def test_pusta_lista(self):
        self.assertEqual(G.przypisz_kanaly_pliku([]), {})


class TestNazwaKolumnyWPliku(unittest.TestCase):

    MAPA = {101: "001"}

    def nazwa(self, n):
        return G._nazwa_kolumny_w_pliku(n, self.MAPA)

    def test_odczyt_kanalu(self):
        self.assertEqual(self.nazwa("Ch101"), "Ch001")

    def test_temperatura_kanalu(self):
        self.assertEqual(self.nazwa("tempCh101"), "tempCh001")

    def test_rozrzut_kanalu(self):
        self.assertEqual(self.nazwa("roztempCh101"), "roztempCh001")

    def test_kolumny_bez_kanalu_bez_zmian(self):
        for n in ("Data Czas", "Tzadana", "tdp", "roztdp(15min)"):
            with self.subTest(kolumna=n):
                self.assertEqual(self.nazwa(n), n)

    def test_kanal_bez_przypisania_bez_zmian(self):
        self.assertEqual(self.nazwa("Ch103"), "Ch103")


class TestParserNaPliku(unittest.TestCase):
    """Pelny odczyt — tak jak robi to krok 2."""

    @classmethod
    def setUpClass(cls):
        cls.folder = nowa_piaskownica("kanaly_cc04")

    def kolumna(self, rows, nazwa, wiersz=0):
        return rows[wiersz][G.CC04_KOLUMNY.index(nazwa)]

    def test_plik_z_jednym_czujnikiem_zostaje_cc04(self):
        """Uzytkownik: to plik CC-04, tylko z jednym czujnikiem."""
        p = plik(os.path.join(self.folder, "typ.txt"), METADANE_221,
                 NAGLOWEK_1_KANAL, ["2026-10-03 16:07:27;23;0;21.8;48.3;"
                                    "108.6110;brak;brak;brak;22.1830;brak;brak"])
        self.assertEqual(G.detect_file_type(G.open_txt(p)), "CC04")

    def test_rezystancja_trafia_do_ukladu(self):
        p = plik(os.path.join(self.folder, "rez.txt"), METADANE_221,
                 NAGLOWEK_1_KANAL, ["2026-10-03 16:07:27;23;0;21.8;48.3;"
                                    "108.6110;brak;brak;brak;22.1830;brak;brak"])
        _nazwy, rows = G.parse_txt_cc04(p)
        self.assertEqual(self.kolumna(rows, "Ch101"), "108.6110")

    def test_temperatura_czujnika_trafia_do_ukladu(self):
        p = plik(os.path.join(self.folder, "temp.txt"), METADANE_221,
                 NAGLOWEK_1_KANAL, ["2026-10-03 16:07:27;23;0;21.8;48.3;"
                                    "108.6110;brak;brak;brak;22.1830;brak;brak"])
        _nazwy, rows = G.parse_txt_cc04(p)
        self.assertEqual(self.kolumna(rows, "tempCh101"), "22.1830")

    def test_nazwa_czujnika_przypisana(self):
        p = plik(os.path.join(self.folder, "nazwa.txt"), METADANE_221,
                 NAGLOWEK_1_KANAL, ["2026-10-03 16:07:27;23;0;21.8;48.3;"
                                    "108.6110;brak;brak;brak;22.1830;brak;brak"])
        nazwy, _rows = G.parse_txt_cc04(p)
        self.assertEqual(nazwy[0], "Pt100-31")

    def test_plik_standardowy_bez_zmian(self):
        """Regresja: zwykly plik CC-04 z kanalami 101..108."""
        meta = list(METADANE_221)
        meta[7] = "Czujnik wzorcowy: Pt100-09; Wejscie pomiarowe Ch: 101"
        kanaly = list(range(101, 109))
        naglowek = ("Data Czas;Tzadana;RHzadana;Todczytana;RHodczytana;"
                    + ";".join(f"Ch{c}" for c in kanaly) + ";tdp;thigro;%rh;"
                    + ";".join(f"tempCh{c}" for c in kanaly) + ";roztdp(15min);"
                    + ";".join(f"roztempCh{c}" for c in kanaly))
        wartosci = (["2026-10-03 16:07:27", "23", "0", "21.8", "48.3"]
                    + [f"10{i}.5" for i in range(8)] + ["5.1", "22.0", "45.0"]
                    + [f"2{i}.1" for i in range(8)] + ["0.05"]
                    + [f"0.0{i}" for i in range(8)])
        p = plik(os.path.join(self.folder, "standard.txt"), meta, naglowek,
                 [";".join(wartosci)])
        nazwy, rows = G.parse_txt_cc04(p)
        self.assertEqual(self.kolumna(rows, "Ch101"), "100.5")
        self.assertEqual(self.kolumna(rows, "Ch108"), "107.5")
        self.assertEqual(self.kolumna(rows, "tempCh103"), "22.1")
        self.assertEqual(nazwy[0], "Pt100-09")


class TestNazwyCzujnikowNaSlotach(unittest.TestCase):
    """Naglowek obserwacji ma mowic 'Pt100-31', a nie 'Ch101'."""

    def test_czujnik_z_kanalu_001_na_slocie_101(self):
        linie = METADANE_221 + [NAGLOWEK_1_KANAL]
        self.assertEqual(G.mapa_slot_pt(linie).get(101), "Pt100-31")

    def test_naglowek_kolumny_z_nazwa_czujnika(self):
        linie = METADANE_221 + [NAGLOWEK_1_KANAL]
        naglowki = G._naglowki_cc04(G.mapa_slot_pt(linie))
        self.assertIn("Wskazania multimetru Pt100-31", naglowki)

    def test_plik_standardowy_bez_zmian(self):
        linie = ["Czujnik wzorcowy: Pt100-09; Wejscie pomiarowe Ch: 101",
                 "Czujnik wzorcowy: Pt100-13; Wejscie pomiarowe Ch: 103",
                 "Data Czas;" + ";".join(f"Ch{c}" for c in range(101, 109))]
        mapa = G.mapa_slot_pt(linie)
        self.assertEqual((mapa.get(101), mapa.get(103)), ("Pt100-09", "Pt100-13"))


class TestPomiarBezHigrometru(unittest.TestCase):
    """
    Zlecenie 221 — pomiar BEZ higrometru wzorcowego (tdp, %rh, roztdp: same 'brak').
    Punkt -10 °C mial nastawe RH 30 %, a komora ponizej zera wilgotnosci nie
    reguluje (~68 %). Porownanie nastawy z odczytem odrzucalo przez to idealnie
    stabilny punkt temperatury jako 'przejscie/suszenie'.

    Wilgotnosci bez higrometru wzorcowego nie da sie wzorcowac, wiec segment
    traktujemy jak punkt samej temperatury. Wyjatek — strefa suszenia/postoju:
    tam rozjazd RH nadal odrzuca segment (inaczej 25-godzinny postoj komory po
    pomiarze stawal sie punktem).
    """

    def segment(self, t, rh_zad, rh_odcz, z_higrometrem=False, minut=40):
        import datetime
        import openpyxl
        start = datetime.datetime(2026, 10, 4, 1, 20)
        data = []
        for i in range(minut):
            data.append((
                start + datetime.timedelta(minutes=i),
                t, rh_zad,
                0.05 if z_higrometrem else None,     # roztdp
                0.01, None, None, None,              # rozrzut temp. kanalow glownych
                t, rh_odcz,                          # odczyt komory
            ))
        wb = openpyxl.Workbook()
        ws = wb.active
        for k, nazwa in enumerate(G.CC04_KOLUMNY, 1):
            ws.cell(row=1, column=k).value = nazwa
        for r in range(minut):
            ws.cell(row=2 + r, column=1).value = data[r][0]
        self.addCleanup(wb.close)
        import io
        from contextlib import redirect_stdout
        with redirect_stdout(io.StringIO()) as buf:
            wynik = G._process_segment(ws, data, 0, len(data), 1, "CC04")
        self.log = buf.getvalue()
        return wynik

    def test_punkt_minus_10_zostaje(self):
        rep, powod = self.segment(-10.3, 30.0, 67.9)
        self.assertIsNotNone(rep, "punkt -10 °C nie moze zostac pominiety")
        self.assertIsNone(powod, "temperatura stabilna — punkt zielony")

    def test_log_mowi_dlaczego_bez_wilgotnosci(self):
        self.segment(-10.3, 30.0, 67.9)
        self.assertIn("Tylko temperatura", self.log)

    def test_postoj_komory_nadal_odrzucany(self):
        """23 °C / 30 % — strefa suszenia; rozjazd RH dalej znaczy 'to nie punkt'."""
        rep, _powod = self.segment(23.0, 30.0, 90.6)
        self.assertIsNone(rep)

    def test_z_higrometrem_zachowanie_bez_zmian(self):
        """Gdy higrometr jest — rozjazd nastawy RH odrzuca segment jak dotad."""
        rep, _powod = self.segment(-10.3, 30.0, 67.9, z_higrometrem=True)
        self.assertIsNone(rep)


if __name__ == "__main__":
    unittest.main(verbosity=2)
