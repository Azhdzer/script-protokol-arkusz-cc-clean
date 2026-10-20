# -*- coding: utf-8 -*-
"""
Zlecenie 229: termometry testo 175T2 (tylko temperatura), producent z dopiskiem.

    1) Termometr elektroniczny (rejestrator, 11 szt.) typ: testo 175T2,
       • nr fabr.: 40118745, nr ewid.: Q/LOG/30,
       ...
       • nr fabr.: 40119897, nr ewid.: Q/LOG/56
       wytworca: Testo – wzorcowanie czujnika wewnetrznego.

Strona 2 protokolu wychodzila inaczej niz wypelniana recznie:
  * Wytworca: 'Testo – wzorcowanie czujnika wewnetrznego' — a ma byc 'Testo',
    dopisek idzie do kolumny 'Uwagi',
  * Typ: 'testo 175T2' — producent stoi obok, wiec samo '175T2',
  * ostatni podpunkt (bez przecinka) mial nr ewid. 'Q/LOG/56 wytworca: Testo – ...',
  * Rozdzielczosc RH: 0,1 (przepisana z temperatury) — a termometr nie mierzy RH: '-',
  * Skalowanie przetwornika t / RH: puste — ma byc '-'.
Strona 1: przy pomiarze samej temperatury higrometr wzorcowy (D12) i rotametr
(F64) nie braly udzialu — '-'.

Fragmenty tekstu zamiast PDF — PZ zawiera dane zleceniodawcy.
"""

import io
import os
import sys
import unittest
from contextlib import redirect_stdout

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import generuj_obserwacje as G
import pz_dane

TEKST = (
    "Numery zleceń laboratorium: 543/LA/T/2026; 229/LA/TH/2026 \n"
    "Obiekty wzorcowania: \n"
    "1) Termometr elektroniczny (rejestrator, 3 szt.) typ: testo 175T2, \n"
    " nr fabr.: 40118745, nr ewid.: Q/LOG/30, \n"
    " nr fabr.: 40119890, nr ewid.: Q/LOG/54 \n"
    " nr fabr.: 40119897, nr ewid.: Q/LOG/56 \n"
    "wytwórca: Testo – wzorcowanie czujnika wewnętrznego. \n"
    "2) Termometr elektroniczny (rejestrator) typ: testo 175T2, nr fabr.: 40118611, "
    "nr ewid.: Q/LOG/40,  \nwytwórca: Testo – wzorcowanie czujnika wewnętrznego. \n"
    "3) Termometr elektroniczny (rejestrator) typ: testo 175T2, nr fabr.: 40118611, "
    "nr ewid.: Q/LOG/40  oraz czujnika \ntemperatury typ: 0572 1001, nr: 8611, "
    "wytwórca: Testo – wzorcowanie czujnika zewnętrznego. \n"
    "Metoda wzorcowania:  \nWzorcowanie w punktach pomiarowych: \n"
    "1) ÷ 2) Metoda porównawcza w zakresie akredytacji AP 074 w komorze klimatycznej, "
    "zgodnie z instrukcją ILAJ 5.4/11 Wzorcowanie termohigrometrów. \n"
    "3) Metoda porównawcza w zakresie akredytacji AP 074, w termostacie cieczowym, "
    "zgodnie z instrukcją ILAJ 5.4/3. \n"
)


def przyrzady():
    return pz_dane.parsuj_tekst(TEKST)


class TestParsowanie(unittest.TestCase):

    def setUp(self):
        self.p = przyrzady()
        self.komora = [p for p in self.p if p.komora]

    def test_wszystkie_przyrzady(self):
        self.assertEqual([p.nr_fabr for p in self.komora],
                         ["40118745", "40119890", "40119897", "40118611"])

    def test_producent_bez_dopisku(self):
        for p in self.p:
            with self.subTest(fabr=p.nr_fabr):
                self.assertEqual(p.wytworca, "Testo")

    def test_dopisek_w_uwagach(self):
        self.assertEqual(self.komora[0].uwagi, "– wzorcowanie czujnika wewnętrznego")

    def test_dopisek_czujnika_zewnetrznego(self):
        zewn = [p for p in self.p if not p.komora][0]
        self.assertEqual(zewn.czuj_wytworca, "Testo")
        self.assertEqual(zewn.uwagi, "– wzorcowanie czujnika zewnętrznego")

    def test_ostatni_podpunkt_bez_ogona(self):
        """Regresja: 'Q/LOG/56 wytwórca: Testo – wzorcowanie czujnika wewnętrznego'."""
        self.assertEqual([p.nr_ewid for p in self.komora],
                         ["Q/LOG/30", "Q/LOG/54", "Q/LOG/56", "Q/LOG/40"])


class TestPomocnicze(unittest.TestCase):

    def test_typ_bez_producenta(self):
        self.assertEqual(pz_dane.typ_bez_producenta("testo 175T2", "Testo"), "175T2")

    def test_typ_bez_producenta_inny_producent(self):
        self.assertEqual(pz_dane.typ_bez_producenta("HMP110", "Vaisala"), "HMP110")

    def test_typ_rowny_producentowi_zostaje(self):
        self.assertEqual(pz_dane.typ_bez_producenta("TERMIOPLUS", "TERMIOPLUS"),
                         "TERMIOPLUS")

    def test_myslnik_w_nazwie_producenta_zostaje(self):
        self.assertEqual(pz_dane.rozdziel_wytworce("LAB-EL"), ("LAB-EL", ""))

    def test_producent_bez_dopisku_bez_zmian(self):
        self.assertEqual(pz_dane.rozdziel_wytworce("Rotronic"), ("Rotronic", ""))


class TestStrona2(unittest.TestCase):
    """Wiersz Strony 2 dla termometru z PZ 229."""

    def setUp(self):
        mapa = {pz_dane.normalizuj_serial(p.nr_fabr): p for p in przyrzady() if p.komora}
        wb = openpyxl.Workbook()
        self.addCleanup(wb.close)
        self.ws = wb.active
        temps = [round(20.0 + i * 0.1, 1) for i in range(40)]
        with redirect_stdout(io.StringIO()):
            G.wypelnij_strone2_z_pz(self.ws, [("40118745", temps, [])], mapa, [],
                                    nr_zlecenia="229")
        self.w = G.STRONA2_PIERWSZY_WIERSZ

    def k(self, kol):
        return self.ws[f"{kol}{self.w}"].value

    def test_jak_wypelnione_recznie(self):
        self.assertEqual(self.k("B"), "Testo")
        self.assertEqual(self.k("D"), "175T2")
        self.assertEqual(self.k("E"), "40118745")
        self.assertEqual(self.k("F"), "Q/LOG/30")
        self.assertEqual(self.k("K"), 0.1)
        self.assertEqual(self.k("L"), "-")
        self.assertEqual(self.k("M"), "-")
        self.assertEqual(self.k("N"), "-")
        self.assertEqual(self.k("O"), "229")
        self.assertEqual(self.k("P"), "– wzorcowanie czujnika wewnętrznego")

    def test_wpisane_skalowanie_nie_jest_nadpisywane(self):
        self.ws[f"M{self.w}"] = "4-20 mA"
        G._skalowanie_nie_dotyczy(self.ws, self.w)
        self.assertEqual(self.k("M"), "4-20 mA")


class TestBrakujacyPlik(unittest.TestCase):

    def test_przyrzad_z_pz_bez_pliku_zgloszony(self):
        with redirect_stdout(io.StringIO()) as buf:
            brak = G._ostrzez_o_brakujacych_plikach(
                przyrzady(), [("40118745", [], []), ("40119890", [], []),
                              ("40119897", [], [])])
        self.assertEqual([p.nr_fabr for p in brak], ["40118611"])
        self.assertIn("40118611", buf.getvalue())

    def test_termostat_nie_zglaszany(self):
        """Pozycja 3 (termostat) nie nalezy do tego protokolu."""
        with redirect_stdout(io.StringIO()):
            brak = G._ostrzez_o_brakujacych_plikach(
                przyrzady(), [("40118745", [], []), ("40119890", [], []),
                              ("40119897", [], []), ("40118611", [], [])])
        self.assertEqual(brak, [])


class TestKrok3(unittest.TestCase):
    """'-' w kolumnie L nie moze trafic do H55 kopii — psuloby formuly RH."""

    def test_liczba(self):
        import generuj_arkusze as A
        self.assertEqual(A._liczba_lub_none(0.1), 0.1)
        self.assertEqual(A._liczba_lub_none("0,01"), 0.01)

    def test_kreska_i_pusto(self):
        import generuj_arkusze as A
        for v in ("-", "", None, True):
            with self.subTest(v=v):
                self.assertIsNone(A._liczba_lub_none(v))


class TestStrona1(unittest.TestCase):

    def test_komorki_higrometru_i_rotametru(self):
        self.assertEqual(G.STRONA1_TYLKO_RH, ((12, 4), (64, 6)))


if __name__ == "__main__":
    unittest.main(verbosity=2)
