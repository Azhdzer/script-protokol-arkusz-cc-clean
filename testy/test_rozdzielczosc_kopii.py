# -*- coding: utf-8 -*-
"""
Testy przenoszenia rozdzielczosci odczytu z protokolu do kopii arkusza obliczen.

Zgloszenie: w protokole (Strona 2) obie rozdzielczosci byly ustawione na 0,01,
a w kopii arkusza obliczeniowego H55 pokazywalo 0,1.

Przyczyna: krok 3 czytal ze Strony 2 tylko kolumne K (rozdzielczosc temperatury)
i wpisywal ja do H57. Kolumna L (rozdzielczosc wilgotnosci) nie byla czytana
w ogole, wiec H55 zostawalo takie, jak w szablonie arkusza — 0,1. A od H55
zalezy budzet niepewnosci wilgotnosci, wiec blad szedl dalej do swiadectwa.

Mapowanie:
    Strona 2, kolumna K  ->  H55? NIE. -> H57  (temperatura)
    Strona 2, kolumna L  ->  H55          (wilgotnosc wzgledna)

Kolejnosc jest myląca (K idzie nizej niz L), dlatego ponizszy test pilnuje, ktory
wiersz szablonu jest ktory — gdyby szablon kiedys przestawil wiersze, zamiana
H55/H57 byla by cicha i widoczna dopiero w liczbach niepewnosci.

Sprawdzenie na gotowej kopii robi test_obieg.py
(test_3_rozdzielczosc_z_protokolu_trafia_do_kopii) — wymaga Excela.
"""

import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from wspolne import KORZEN, SZABLON_ARKUSZA

WIERSZ_RH = 55
WIERSZ_T = 57


class TestWierszeSzablonu(unittest.TestCase):
    """Ktora komorka szablonu jest wilgotnoscia, a ktora temperatura."""

    @classmethod
    def setUpClass(cls):
        sciezka = os.path.join(KORZEN, SZABLON_ARKUSZA)
        if not os.path.exists(sciezka):
            raise unittest.SkipTest(f"brak szablonu: {SZABLON_ARKUSZA}")
        wb = openpyxl.load_workbook(sciezka, read_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            cls.opisy = {r: cls._etykieta(ws, r) for r in (WIERSZ_RH, WIERSZ_T)}
        finally:
            wb.close()

    @staticmethod
    def _etykieta(ws, wiersz):
        for kol in range(1, 8):
            v = ws.cell(row=wiersz, column=kol).value
            if v not in (None, ""):
                return str(v).lower()
        return ""

    def test_h55_dotyczy_wilgotnosci(self):
        self.assertIn("wilgotn", self.opisy[WIERSZ_RH])

    def test_h57_dotyczy_temperatury(self):
        self.assertIn("temperatur", self.opisy[WIERSZ_T])

    def test_oba_wiersze_mowia_o_rozdzielczosci(self):
        for wiersz, opis in self.opisy.items():
            with self.subTest(wiersz=wiersz):
                self.assertIn("rozdzielczo", opis)

    def test_wiersze_nie_sa_zamienione(self):
        """Gdyby szablon przestawil wiersze, H55/H57 poszlyby na krzyz."""
        self.assertNotIn("temperatur", self.opisy[WIERSZ_RH])
        self.assertNotIn("wilgotn", self.opisy[WIERSZ_T])


class TestKolumnyStrony2(unittest.TestCase):
    """Naglowki Strony 2 protokolu — zrodlo obu rozdzielczosci."""

    PROTOKOL = "xxx_LA_TH_2026 - protokół CC.xlsx"
    WIERSZ_NAGLOWKA = 10

    @classmethod
    def setUpClass(cls):
        sciezka = os.path.join(KORZEN, cls.PROTOKOL)
        if not os.path.exists(sciezka):
            raise unittest.SkipTest(f"brak protokolu wzorcowego: {cls.PROTOKOL}")
        wb = openpyxl.load_workbook(sciezka, read_only=True)
        try:
            ws = wb["Strona 2"]
            cls.naglowki = {
                k: str(ws.cell(row=cls.WIERSZ_NAGLOWKA, column=k).value or "").lower()
                for k in (11, 12)
            }
        finally:
            wb.close()

    def test_kolumna_k_to_temperatura(self):
        self.assertIn("t °c", self.naglowki[11])

    def test_kolumna_l_to_wilgotnosc(self):
        self.assertIn("rh", self.naglowki[12])

    def test_obie_kolumny_to_rozdzielczosc(self):
        for kol, opis in self.naglowki.items():
            with self.subTest(kolumna=kol):
                self.assertIn("rozdzielczo", opis)


if __name__ == "__main__":
    unittest.main(verbosity=2)
