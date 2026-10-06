# -*- coding: utf-8 -*-
"""
Testy opracowania plikow kroku 1 przez krok 2: nazwa zlecenia i podpis.

Krok 1 zapisuje '<serial>_wynik.xlsx' — zlecenia jeszcze nie zna. Krok 2 zna je
z pliku TXT multimetru, wiec po zaznaczeniu punktow:

    10705098_wynik.xlsx        ->  221_LA_TH_2026_10705098_opracowanie danych.xlsx
    zestawienie_pomiarow.xlsx  ->  221_LA_TH_2026_zestawienie_pomiarow.xlsx

i dopisuje z prawej strony tabeli podpis:

    Opracowal:          Data:
    Artsiom Azhdzer     06.10.2026

Pulapki, ktorych pilnuja te testy:
  * Ponowny przebieg kroku 2 czyta juz opracowane pliki — numer fabryczny musi
    sie dac wyciagnac takze z nowej nazwy.
  * Ponowny krok 1 kladzie swiezy '<serial>_wynik.xlsx' OBOK opracowanego — ten
    sam przyrzad nie moze trafic do protokolu dwa razy.
  * Znacznik punktow liczyl kolumne od ws.max_column; po opracowaniu stoja tam
    'Nr punktu' i podpis, wiec przy kazdym przebiegu znacznik wedrowalby w prawo.
  * Pliki NIEDOPASOWANE do punktow (inne wzorcowanie) zostaja nietkniete.
"""

import datetime
import io
import os
import sys
import unittest
from contextlib import redirect_stdout

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

import generuj_obserwacje as G
from wspolne import nowa_piaskownica


def plik_wyniku(sciezka, naglowki=("Czas", "Temperatura [°C]", "Wilgotność [%RH]"),
                wierszy=12):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(naglowki))
    for i in range(wierszy):
        ws.append([datetime.datetime(2026, 10, 3, 16, i), 21.8 + i * 0.01, 48.0][:len(naglowki)])
    wb.save(sciezka)
    return sciezka


class ZFolderemWynikow(unittest.TestCase):
    """Przekierowuje WYNIKI_FOLDER do piaskownicy na czas testu."""

    NAZWA = "opracowanie_wynikow"

    def setUp(self):
        self.folder = nowa_piaskownica(f"{self.NAZWA}_{self._testMethodName}")
        self.wyniki = os.path.join(self.folder, "wyniki")
        os.makedirs(self.wyniki, exist_ok=True)
        self._stary = G.WYNIKI_FOLDER
        G.WYNIKI_FOLDER = self.wyniki
        del G._PLIKI_ZAZNACZONE[:]

    def tearDown(self):
        G.WYNIKI_FOLDER = self._stary
        del G._PLIKI_ZAZNACZONE[:]

    def sciezka(self, nazwa):
        return os.path.join(self.wyniki, nazwa)

    def pliki(self):
        return sorted(os.listdir(self.wyniki))


class TestNazwy(unittest.TestCase):

    def test_przedrostek_jak_w_nazwie_protokolu(self):
        self.assertEqual(G.prefiks_zlecenia("221", "CC04"), "221_LA_TH_2026")
        self.assertEqual(G.prefiks_zlecenia("213", "CC"), "213_LA_TH_2026")

    def test_nazwa_pliku_przyrzadu(self):
        self.assertEqual(G.nazwa_pliku_opracowania("221_LA_TH_2026", "10705098"),
                         "221_LA_TH_2026_10705098_opracowanie danych.xlsx")

    def test_nazwa_zestawienia(self):
        self.assertEqual(G.nazwa_pliku_zestawienia("221_LA_TH_2026"),
                         "221_LA_TH_2026_zestawienie_pomiarow.xlsx")

    def test_serial_z_opracowanego_pliku(self):
        """Ponowny krok 2 czyta juz opracowane pliki."""
        self.assertEqual(
            G._serial_z_wyniku("221_LA_TH_2026_10705098_opracowanie danych.xlsx"),
            "10705098")

    def test_serial_ze_swiezego_pliku_bez_zmian(self):
        self.assertEqual(G._serial_z_wyniku("10705098_wynik.xlsx"), "10705098")

    def test_zestawienie_rozpoznawane_w_obu_postaciach(self):
        for n in ("zestawienie_pomiarow.xlsx", "221_LA_TH_2026_zestawienie_pomiarow.xlsx"):
            with self.subTest(nazwa=n):
                self.assertTrue(G._RE_PLIK_ZESTAWIENIA.search(n))


class TestListaPlikow(ZFolderemWynikow):

    def test_opracowane_zestawienie_nie_jest_przyrzadem(self):
        plik_wyniku(self.sciezka("221_LA_TH_2026_zestawienie_pomiarow.xlsx"))
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        self.assertEqual(G._pliki_wynikow(), ["10705098_wynik.xlsx"])

    def ustaw_czas(self, nazwa, sekundy_temu):
        t = datetime.datetime.now().timestamp() - sekundy_temu
        os.utime(self.sciezka(nazwa), (t, t))

    def test_swiezy_wynik_wypiera_starszy_opracowany(self):
        """Ponowny krok 1 bez PZ — przyrzad nie moze trafic do protokolu dwa razy."""
        plik_wyniku(self.sciezka("221_LA_TH_2026_10705098_opracowanie danych.xlsx"))
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        self.ustaw_czas("221_LA_TH_2026_10705098_opracowanie danych.xlsx", 3600)
        self.assertEqual(G._pliki_wynikow(), ["10705098_wynik.xlsx"])

    def test_swiezy_opracowany_wypiera_starszy_wynik(self):
        """Krok 1 z PZ zapisuje od razu nazwe zlecenia — stary '_wynik' nie liczy sie."""
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        plik_wyniku(self.sciezka("221_LA_TH_2026_10705098_opracowanie danych.xlsx"))
        self.ustaw_czas("10705098_wynik.xlsx", 3600)
        self.assertEqual(G._pliki_wynikow(),
                         ["221_LA_TH_2026_10705098_opracowanie danych.xlsx"])

    def test_sam_opracowany_jest_czytany(self):
        """Ponowny krok 2 bez kroku 1."""
        plik_wyniku(self.sciezka("221_LA_TH_2026_10705098_opracowanie danych.xlsx"))
        self.assertEqual(G._pliki_wynikow(),
                         ["221_LA_TH_2026_10705098_opracowanie danych.xlsx"])

    def test_opracowany_innego_przyrzadu_zostaje(self):
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        plik_wyniku(self.sciezka("213_LA_TH_2026_10861648_opracowanie danych.xlsx"))
        self.assertEqual(len(G._pliki_wynikow()), 2)


class TestSzerokoscTabeli(unittest.TestCase):

    def test_liczy_tylko_kolumny_danych(self):
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Czas", "Temperatura [°C]", "Wilgotność [%RH]", None, "Nr punktu",
                   None, "Opracował:", "Data:"])
        self.assertEqual(G._szerokosc_tabeli_wynikow(ws), 3)

    def test_przyrzad_tylko_temperatura(self):
        wb = openpyxl.Workbook(); ws = wb.active
        ws.append(["Czas", "Temperatura [°C]"])
        self.assertEqual(G._szerokosc_tabeli_wynikow(ws), 2)


class TestOpracowanie(ZFolderemWynikow):

    DATA = datetime.date(2026, 10, 6)

    def opracuj(self, prefiks="221_LA_TH_2026"):
        with redirect_stdout(io.StringIO()):
            return G.opracuj_pliki_wynikow(prefiks, podpis="Artsiom Azhdzer", data=self.DATA)

    def oznacz(self, nazwa, punkty=None):
        with redirect_stdout(io.StringIO()):
            G._oznacz_wyniki_xlsx(self.sciezka(nazwa), punkty or {0: {3, 4, 5}})

    def arkusz(self, nazwa):
        wb = openpyxl.load_workbook(self.sciezka(nazwa))
        self.addCleanup(wb.close)
        return wb.active

    def test_zaznaczony_plik_dostaje_nazwe_zlecenia(self):
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        self.oznacz("10705098_wynik.xlsx")
        self.opracuj()
        self.assertEqual(self.pliki(), ["221_LA_TH_2026_10705098_opracowanie danych.xlsx"])

    def test_podpis_i_data_w_pliku(self):
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        self.oznacz("10705098_wynik.xlsx")
        self.opracuj()
        ws = self.arkusz("221_LA_TH_2026_10705098_opracowanie danych.xlsx")
        # dane A-C, odstep D, 'Nr punktu' E, odstep F, podpis G-H
        self.assertEqual((ws["G1"].value, ws["H1"].value), ("Opracował:", "Data:"))
        self.assertEqual(ws["G2"].value, "Artsiom Azhdzer")
        self.assertEqual(ws["H2"].value.date(), self.DATA)
        self.assertEqual(ws["H2"].number_format, "dd.mm.yyyy")

    def test_podpis_nie_zaslania_danych_ani_znacznika(self):
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        self.oznacz("10705098_wynik.xlsx")
        self.opracuj()
        ws = self.arkusz("221_LA_TH_2026_10705098_opracowanie danych.xlsx")
        self.assertEqual(ws["C1"].value, "Wilgotność [%RH]")
        self.assertEqual(ws["E1"].value, "Nr punktu")
        self.assertEqual(ws["C2"].value, 48.0)

    def test_niezaznaczony_plik_zostaje_nietkniety(self):
        """Inne wzorcowanie — ani nowej nazwy, ani podpisu."""
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        plik_wyniku(self.sciezka("99999999_wynik.xlsx"))
        self.oznacz("10705098_wynik.xlsx")
        self.opracuj()
        self.assertIn("99999999_wynik.xlsx", self.pliki())
        self.assertIsNone(self.arkusz("99999999_wynik.xlsx")["G1"].value)

    def test_zestawienie_tez_opracowane(self):
        plik_wyniku(self.sciezka("zestawienie_pomiarow.xlsx"),
                    naglowki=("Czas", "Temp 10705098", "Wilg 10705098"))
        self.opracuj()
        self.assertEqual(self.pliki(), ["221_LA_TH_2026_zestawienie_pomiarow.xlsx"])
        ws = self.arkusz("221_LA_TH_2026_zestawienie_pomiarow.xlsx")
        self.assertEqual(ws["G1"].value, "Opracował:")

    def test_ponowny_przebieg_nie_przesuwa_znacznika_ani_podpisu(self):
        """Po opracowaniu w arkuszu stoja 'Nr punktu' i podpis — kolumny maja zostac."""
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        self.oznacz("10705098_wynik.xlsx")
        self.opracuj()
        nazwa = "221_LA_TH_2026_10705098_opracowanie danych.xlsx"
        del G._PLIKI_ZAZNACZONE[:]
        self.oznacz(nazwa)
        self.opracuj()
        ws = self.arkusz(nazwa)
        self.assertEqual(ws["E1"].value, "Nr punktu")
        self.assertEqual(ws["G1"].value, "Opracował:")
        self.assertIsNone(ws["I1"].value, "znacznik lub podpis powedrowal w prawo")
        self.assertEqual(self.pliki(), [nazwa])

    def test_swiezy_plik_zastepuje_opracowany(self):
        """Ponowny krok 1 + krok 2: zostaje jeden plik, ze swiezymi danymi."""
        nazwa = "221_LA_TH_2026_10705098_opracowanie danych.xlsx"
        plik_wyniku(self.sciezka(nazwa), wierszy=3)
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"), wierszy=12)
        self.oznacz("10705098_wynik.xlsx")
        self.opracuj()
        self.assertEqual(self.pliki(), [nazwa])
        self.assertEqual(self.arkusz(nazwa).max_row, 13)   # naglowek + 12 swiezych


class TestZlecenieZPZ(ZFolderemWynikow):
    """
    Numer zlecenia przyrzadu z PZ — szukany po nr fabrycznym albo ewidencyjnym.
    Jeden wsad komory obejmuje czasem kilka zlecen: krok 3 nazywa kopie wg
    zlecenia PRZYRZADU, wiec plik opracowania musi nosic ten sam numer.
    """

    def mapa(self, nr_zlecenia="195", nr_fabr="10705098", nr_ewid="Onset-06"):
        import pz_dane
        p = pz_dane.PZPrzyrzad(nr_zlecenia=nr_zlecenia, nr_fabr=nr_fabr,
                               nr_ewid=nr_ewid, komora=True)
        return {pz_dane.normalizuj_serial(nr_fabr): p,
                pz_dane.normalizuj_serial(nr_ewid): p}

    def opracuj(self, pz_mapa):
        with redirect_stdout(io.StringIO()):
            G.opracuj_pliki_wynikow("221_LA_TH_2026", podpis="X",
                                    data=datetime.date(2026, 10, 6),
                                    pz_mapa=pz_mapa, obs_type="CC04")

    def test_krok2_bierze_zlecenie_przyrzadu_z_pz(self):
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        with redirect_stdout(io.StringIO()):
            G._oznacz_wyniki_xlsx(self.sciezka("10705098_wynik.xlsx"), {0: {3}})
        self.opracuj(self.mapa("195"))
        self.assertEqual(self.pliki(), ["195_LA_TH_2026_10705098_opracowanie danych.xlsx"])

    def test_krok2_bez_przyrzadu_w_pz_numer_pomiaru(self):
        plik_wyniku(self.sciezka("10705098_wynik.xlsx"))
        with redirect_stdout(io.StringIO()):
            G._oznacz_wyniki_xlsx(self.sciezka("10705098_wynik.xlsx"), {0: {3}})
        self.opracuj(self.mapa("195", nr_fabr="99999999", nr_ewid="INNY"))
        self.assertEqual(self.pliki(), ["221_LA_TH_2026_10705098_opracowanie danych.xlsx"])

    def test_zestawienie_zawsze_z_numerem_pomiaru(self):
        """Zestawienie dotyczy calego wsadu, nie jednego przyrzadu."""
        plik_wyniku(self.sciezka("zestawienie_pomiarow.xlsx"))
        self.opracuj(self.mapa("195"))
        self.assertEqual(self.pliki(), ["221_LA_TH_2026_zestawienie_pomiarow.xlsx"])


class TestKrok1NazwaZPZ(unittest.TestCase):
    """Krok 1 nadaje nazwe zlecenia od razu, gdy przyrzad jest w PZ."""

    def setUp(self):
        import analizuj_excele as A
        import pz_dane
        self.A, self.P = A, pz_dane
        self.folder = nowa_piaskownica(f"krok1_pz_{self._testMethodName}")
        self._stara = A._PZ_MAPA
        p = pz_dane.PZPrzyrzad(nr_zlecenia="221", nr_fabr="10705098",
                               nr_ewid="Onset-06", komora=True)
        A._PZ_MAPA = {pz_dane.normalizuj_serial("10705098"): p,
                      pz_dane.normalizuj_serial("Onset-06"): p}

    def tearDown(self):
        self.A._PZ_MAPA = self._stara

    def nazwa(self, stem, suffix=""):
        with redirect_stdout(io.StringIO()):
            return self.A.sciezka_wyniku(self.folder, stem, suffix).name

    def test_przyrzad_w_pz_dostaje_nazwe_zlecenia(self):
        self.assertEqual(self.nazwa("10705098"),
                         "221_LA_TH_2026_10705098_opracowanie danych.xlsx")

    def test_dopasowanie_po_numerze_ewidencyjnym(self):
        """Pliki loggerow bywaja nazwane nr ewidencyjnym zamiast fabrycznego."""
        self.assertEqual(self.nazwa("Onset-06"),
                         "221_LA_TH_2026_Onset-06_opracowanie danych.xlsx")

    def test_data_w_nazwie_pliku_loggera(self):
        self.assertEqual(self.nazwa("2026-10-03-16-07-27 10705098"),
                         "221_LA_TH_2026_10705098_opracowanie danych.xlsx")

    def test_przyrzad_spoza_pz_po_staremu(self):
        self.assertEqual(self.nazwa("99999999"), "99999999_wynik.xlsx")

    def test_czesc_pliku_po_staremu(self):
        """Osobny arkusz/kanal (suffix) — bez zmian, jak dotad."""
        self.assertEqual(self.nazwa("10705098", "_Arkusz2"), "10705098_Arkusz2_wynik.xlsx")

    def test_stary_wynik_tego_zrodla_usuniety(self):
        """Inaczej przyrzad lezalby w 'wyniki' dwa razy."""
        stary = os.path.join(self.folder, "10705098_wynik.xlsx")
        plik_wyniku(stary)
        self.nazwa("10705098")
        self.assertFalse(os.path.exists(stary))

    def test_bez_pz_nic_nie_jest_usuwane(self):
        self.A._PZ_MAPA = {}
        stary = os.path.join(self.folder, "10705098_wynik.xlsx")
        plik_wyniku(stary)
        self.assertEqual(self.nazwa("10705098"), "10705098_wynik.xlsx")
        self.assertTrue(os.path.exists(stary))


if __name__ == "__main__":
    unittest.main(verbosity=2)
