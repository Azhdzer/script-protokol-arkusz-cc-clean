# -*- coding: utf-8 -*-
"""
wypelnij_warunki.py — krok 4: warunki srodowiskowe w protokole.

Wpisuje temperature i wilgotnosc pomieszczenia (Strona 3, kolumny F/G
'Warunki srodowiskowe') z pliku czujnika 'Pom. nr 9 (MX1101-02) - 2026.xlsx'
dla KAZDEGO punktu protokolu — na chwile startu i konca punktu.

Dlaczego osobny krok
--------------------
Krok 3 tez to robi, ale tylko przy okazji tworzenia kopii arkusza: wypelnia F/G
dla blokow, ktore trafily do jakiejs kopii. Punkt wyszarzony dla wszystkich
przyrzadow (np. spoza ich zamowienia) do zadnej kopii nie trafia, wiec jego
warunki zostawaly puste. Ten krok pracuje wprost na protokole, niezaleznie od
kopii, Worda i Excela (nie uruchamia Excela w ogole).

Uruchomiony PRZED krokiem 3 tez sie przydaje — krok 3 widzi wtedy gotowe F/G
('F/G juz wypelnione') i nie musi ich szukac.

Wyszukiwanie w pliku czujnika jest DOKLADNIE to samo, co w kroku 3 (te same
funkcje z generuj_arkusze): najblizszy rekord w czasie, ale nie dalej niz
tolerancja z panelu — inaczej komorka zostaje pusta do recznego uzupelnienia.
"""

import datetime
import os
import sys

import openpyxl

import cc_config as C
import cc_widok
# Te same funkcje i stale co w kroku 3 — wynik ma byc identyczny.
import generuj_arkusze as A

FOLDER     = A.FOLDER
PROTOKOL   = C.tekst("SR_PROTOKOL", "")
NADPISUJ   = C.flaga("SR_NADPISUJ", False)
TAB_RATIO  = C.liczba("GEN_TAB_RATIO", cc_widok.TAB_RATIO_DOMYSLNY)

# Kolumny warunkow srodowiskowych w bloku punktu (Strona 3).
KOL_T, KOL_RH = 6, 7       # F, G
KOL_CZASU = 5              # E: r0+1 data, r0+2 start, r0+3 koniec


def bloki_strony3(ws):
    """
    Bloki punktow ze Strony 3: (nr_bloku, wiersz_poczatku, data, start, koniec).
    Blok istnieje, dopoki w kolumnie A jest numer punktu — jak w kroku 3.
    """
    nr = 0
    while True:
        r0 = A.START_ROW_S3 + nr * A.BLOK_S3
        if ws.cell(row=r0, column=1).value in (None, ""):
            return
        yield (nr, r0,
               ws.cell(row=r0 + 1, column=KOL_CZASU).value,
               ws.cell(row=r0 + 2, column=KOL_CZASU).value,
               ws.cell(row=r0 + 3, column=KOL_CZASU).value)
        nr += 1


def wypelnij_protokol(sciezka, nadpisuj=False, nr_pom=None, model=None):
    """
    Wypelnia F/G w Stronie 3 protokolu. Zwraca slownik z podsumowaniem:
      wpisane    — ile komorek par (T, RH) wpisano,
      pominiete  — ile bylo juz wypelnionych (bez nadpisywania),
      brak       — ile momentow nie mialo rekordu czujnika w tolerancji,
      bloki      — ile blokow obejrzano.
    """
    nr_pom = A.NR_POMIESZCZENIA if nr_pom is None else nr_pom
    model = A.MODEL_CZUJNIKA if model is None else model
    wynik = {"wpisane": 0, "pominiete": 0, "brak": 0, "bloki": 0}

    wb = openpyxl.load_workbook(sciezka)
    if A.ARKUSZ_STRONA3 not in wb.sheetnames:
        wb.close()
        raise ValueError(f"W pliku nie ma arkusza '{A.ARKUSZ_STRONA3}' — "
                         f"to nie wyglada na protokol.")
    ws = wb[A.ARKUSZ_STRONA3]

    cache = {}

    def dane_miesiaca(rok, miesiac):
        if (rok, miesiac) not in cache:
            plik = A._sciezka_pliku_czujnika(nr_pom, model, rok)
            cache[(rok, miesiac)] = A._wczytaj_dane_czujnika_miesiac(plik, rok, miesiac)
            print(f"    plik czujnika: {os.path.basename(plik)}  arkusz "
                  f"{miesiac:02d}'{str(rok)[2:]}  rekordow: {len(cache[(rok, miesiac)])}")
        return cache[(rok, miesiac)]

    for nr, r0, data_val, start_val, koniec_val in bloki_strony3(ws):
        wynik["bloki"] += 1
        data = A._parse_date_s3(data_val)
        if data is None:
            print(f"  punkt {nr + 1}: brak daty w E{r0 + 1} — pomijam")
            continue
        dane = dane_miesiaca(data.year, data.month)
        for czas_val, offset, etykieta in ((start_val, 2, "start"), (koniec_val, 3, "koniec")):
            r = r0 + offset
            czas = A._parse_time_s3(czas_val)
            if czas is None:
                print(f"  punkt {nr + 1} {etykieta}: brak godziny w E{r} — pomijam")
                continue
            kom_t, kom_rh = ws.cell(row=r, column=KOL_T), ws.cell(row=r, column=KOL_RH)
            if not nadpisuj and kom_t.value not in (None, "") and kom_rh.value not in (None, ""):
                wynik["pominiete"] += 1
                print(f"  punkt {nr + 1} {etykieta}: juz wypelnione "
                      f"({kom_t.value} / {kom_rh.value}) — zostawiam")
                continue
            moment = datetime.datetime.combine(data, czas)
            t, rh, odch_s = A._szukaj_th_w_danych_czujnika(dane, moment)
            if t is None:
                wynik["brak"] += 1
                powod = (f"najblizszy rekord {odch_s / 60:.0f} min od punktu, limit "
                         f"{A.TOLERANCJA_CZUJNIKA_MIN} min" if odch_s is not None
                         else "brak danych czujnika w tym miesiacu")
                print(f"  punkt {nr + 1} {etykieta} {moment:%d.%m.%Y %H:%M}: BRAK ({powod}) "
                      f"— F/G zostaja puste")
                continue
            kom_t.value = t
            if rh is not None:
                kom_rh.value = rh
            wynik["wpisane"] += 1
            print(f"  punkt {nr + 1} {etykieta} {moment:%d.%m.%Y %H:%M}: "
                  f"t = {t}  RH = {rh}")

    if wynik["wpisane"]:
        wb.save(sciezka)
        cc_widok.wymus_widok(sciezka, TAB_RATIO, pierwszy=0)
    wb.close()
    return wynik


def main():
    print("=" * 65)
    print("  Warunki srodowiskowe -> Strona 3 protokolu")
    print("=" * 65)
    if not PROTOKOL:
        print("[BŁĄD] Nie wybrano protokolu (krok 4, pole 'Protokol').")
        return 1
    sciezka = PROTOKOL if os.path.isabs(PROTOKOL) else os.path.join(FOLDER, PROTOKOL)
    if not os.path.exists(sciezka):
        print(f"[BŁĄD] Plik protokolu nie istnieje:\n  {sciezka}")
        return 1
    print(f"  Protokol : {os.path.basename(sciezka)}")
    print(f"  Czujnik  : Pom. nr {A.NR_POMIESZCZENIA} ({A.MODEL_CZUJNIKA}), "
          f"tolerancja {A.TOLERANCJA_CZUJNIKA_MIN} min")
    print(f"  Wypelnione komorki: {'nadpisywane' if NADPISUJ else 'zostawiane bez zmian'}")
    print()
    try:
        w = wypelnij_protokol(sciezka, nadpisuj=NADPISUJ)
    except PermissionError:
        print(f"\n[BŁĄD] Protokol jest OTWARTY w Excelu — zamknij go i uruchom krok ponownie.\n"
              f"  {sciezka}")
        return 1
    except ValueError as e:
        print(f"\n[BŁĄD] {e}")
        return 1

    print()
    print("=" * 65)
    print(f"  Punktow: {w['bloki']}   wpisano: {w['wpisane']}   "
          f"juz wypelnione: {w['pominiete']}   bez danych czujnika: {w['brak']}")
    if w["wpisane"]:
        print(f"  Zapisano: {os.path.basename(sciezka)}")
    else:
        print("  Nic nie zmieniono — plik zostal bez zapisu.")
    print("=" * 65)
    return 0


if __name__ == "__main__":
    sys.exit(main())
